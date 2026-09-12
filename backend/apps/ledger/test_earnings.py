from datetime import date, datetime, timedelta, timezone as dt_timezone
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from asgiref.sync import async_to_sync
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from apps.shifts.models import Employee, Organization
from . import earnings
from .models import Batch, EarningsDelivery, Record, Settings, TelegramContact


class EarningsTests(TestCase):
    def setUp(self):
        self.employee = Employee.objects.first()
        self.org = Organization.objects.first()
        self.other = Organization.objects.exclude(pk=self.org.pk).first()
        self.approver = TelegramContact.objects.create(user_id=777, name="Главный")
        Settings.objects.filter(pk=1).update(approver=self.approver, enabled=False)
        self.now = datetime(2026, 5, 1, 5, tzinfo=dt_timezone.utc)

    def row(self, **changes):
        values = dict(kind="oneoff", date=date(2026, 4, 12), employee=self.employee,
                      employee_name="Старое имя", organization=self.org, amount=Decimal("10.10"))
        values.update(changes)
        return Record.objects.create(**values)

    def test_totals_use_identity_saved_amounts_all_organizations_and_review_buckets(self):
        self.row()
        self.row(organization=self.other, amount=Decimal("20.20"), employee_name="Другое имя")
        self.other.is_active = False
        self.other.save()
        self.row(organization=None, amount=Decimal("30.30"))
        self.row(kind="expense", amount=999)
        self.row(deleted_at=timezone.now(), amount=999)
        self.row(date=date(2026, 3, 31), amount=999)
        self.row(date=date(2026, 5, 1), amount=999)
        self.row(review=True, amount=1)
        self.row(error="Нет тарифа", amount=2)
        self.row(employee=None, amount=3)
        self.row(employee=None, source="auto", amount=4)
        self.row(employee=None, source="auto", amount=0, error="Нет тарифа", review=True)
        self.employee.short_name = "Новое имя"
        self.employee.is_active = False
        self.employee.save()
        # Different identities with the same stored name must not be combined.
        second = Employee.objects.exclude(pk=self.employee.pk).first()
        self.row(employee=second, employee_name="Старое имя", amount=5)
        with patch("apps.ledger.engine.calculate", side_effect=AssertionError("Must not recalculate")):
            data = earnings.report_data("2026-04", self.now)
        self.assertEqual(data["total"], "65.60")
        by_id = {r["employee_id"]: r for r in data["employees"]}
        self.assertEqual(by_id[self.employee.pk]["amount"], "60.60")
        self.assertEqual(by_id[self.employee.pk]["employee_name"], "Новое имя")
        self.assertEqual(data["review_total"], "6.00")
        self.assertEqual(data["unassigned_total"], "4.00")
        self.assertEqual(len(data["review"]), 3)
        self.assertEqual(len(data["unassigned"]), 2)

    def test_api_excel_and_empty_month(self):
        self.row()
        self.row(employee=None, amount=3, employee_name="=1+1")
        url = "/api/shifts/ledger/earnings/?month=2026-04"
        with patch("apps.ledger.earnings.timezone.now", return_value=self.now):
            response = self.client.get(url + "&organization=999&limit=1")
            self.assertEqual(response.status_code, 200)
            data = response.json()
            exported = self.client.get(url + "&format=xlsx")
        self.assertIn("no-store", exported["Cache-Control"])
        wb = load_workbook(BytesIO(exported.content))
        self.assertEqual(wb.sheetnames, ["Заработок", "На проверку"])
        self.assertEqual(Decimal(str(wb["Заработок"]["B6"].value)), Decimal(data["total"]))
        self.assertEqual(wb["На проверку"]["C6"].value, "'=1+1")
        self.assertEqual(wb["На проверку"]["D6"].value, 3)
        self.assertEqual(wb["Заработок"]["B6"].number_format, '#,##0.00')
        self.assertEqual(earnings.report_data("2026-03", self.now)["total"], "0.00")
        self.assertEqual(earnings.report_data("2026-05", self.now)["end"], "2026-05-01")
        self.assertEqual(Record.objects.count(), 2)

    def test_invalid_months_and_format(self):
        for month in (None, "2026-13", "2026-4", "0000-01", "2027-01"):
            with self.subTest(month=month), self.assertRaises(ValidationError):
                earnings.report_data(month, self.now)
        with patch("apps.ledger.earnings.timezone.now", return_value=self.now):
            for suffix in ("", "?month=2026-13", "?month=2026-06", "?month=2026-04&format=bad"):
                self.assertEqual(self.client.get("/api/shifts/ledger/earnings/" + suffix).status_code, 400)

    def test_schedule_time_deduplication_snapshot_and_independence(self):
        row = self.row()
        self.assertIsNone(earnings.schedule(self.now - timedelta(seconds=1)))
        item = earnings.schedule(self.now)
        original = bytes(item.artifact)
        row.amount = 100
        row.save()
        self.assertEqual(earnings.schedule(self.now + timedelta(days=4)).pk, item.pk)
        item.refresh_from_db()
        self.assertEqual(bytes(item.artifact), original)
        self.assertEqual(item.snapshot["total"], "10.10")
        self.assertEqual(earnings.report_data("2026-04", self.now)["total"], "100.00")
        self.assertEqual(EarningsDelivery.objects.count(), 1)
        self.assertEqual(Batch.objects.count(), 0)
        with patch("apps.ledger.earnings.timezone.now", return_value=self.now):
            data = self.client.get("/api/shifts/ledger/earnings/?month=2026-04").json()
        self.assertEqual(data["delivery"]["recipient"], 777)
        self.assertNotIn("snapshot", data["delivery"])

    def test_year_boundary_catchup_leap_month_and_settings(self):
        self.assertEqual(earnings.schedule(datetime(2026, 1, 3, tzinfo=dt_timezone.utc)).month, "2025-12")
        self.assertEqual(earnings.report_data("2024-02", self.now)["end"], "2024-02-29")
        Settings.objects.filter(pk=1).update(earnings_enabled=False)
        self.assertIsNone(earnings.schedule(self.now))
        Settings.objects.filter(pk=1).update(earnings_enabled=True, approver=None)
        self.assertIsNone(earnings.schedule(self.now))
        response = self.client.put("/api/shifts/ledger/settings/", '{"earnings_enabled":false}', content_type="application/json")
        self.assertEqual(response.status_code, 200)
        self.assertFalse(response.json()["earnings_enabled"])

    def test_delivery_success_timeout_retry_and_snapshot(self):
        self.row()
        item = earnings.schedule(self.now)
        bot = AsyncMock()
        bot.send_document.side_effect = TimeoutError()
        self.assertTrue(async_to_sync(earnings.deliver_once)(bot))
        item.refresh_from_db()
        self.assertEqual(item.state, "unknown")
        self.assertFalse(async_to_sync(earnings.deliver_once)(bot))
        url = f"/api/shifts/ledger/earnings-deliveries/{item.pk}/retry/"
        self.assertEqual(self.client.post(url, "{}", content_type="application/json").status_code, 400)
        self.assertEqual(self.client.post(url, '{"confirm_duplicate_risk":true}', content_type="application/json").status_code, 200)
        bot.send_document.side_effect = None
        bot.send_document.return_value = SimpleNamespace(message_id=100)
        self.assertTrue(async_to_sync(earnings.deliver_once)(bot))
        item.refresh_from_db()
        self.assertEqual((item.state, item.attempts, item.message_id), ("sent", 2, 100))
        args = bot.send_document.call_args.args
        self.assertEqual(args[0], 777)
        self.assertEqual(args[1].data, bytes(item.artifact))
        self.assertEqual(args[1].filename, "earnings-2026-04.xlsx")
        self.assertFalse(async_to_sync(earnings.deliver_once)(bot))
        self.assertEqual(self.client.post(url, "{}", content_type="application/json").status_code, 400)

    def test_explicit_telegram_rejections_and_worker_crash(self):
        from aiogram.exceptions import TelegramBadRequest, TelegramForbiddenError, TelegramRetryAfter, TelegramUnauthorizedError
        from aiogram.methods import SendDocument
        item = earnings.schedule(self.now)
        method = SendDocument(chat_id=777, document="test")
        errors = [TelegramBadRequest(method=method, message="bad"), TelegramForbiddenError(method=method, message="blocked"),
                  TelegramUnauthorizedError(method=method, message="token"), TelegramRetryAfter(method=method, message="limit", retry_after=12)]
        for error in errors:
            bot = AsyncMock()
            bot.send_document.side_effect = error
            async_to_sync(earnings.deliver_once)(bot)
            item.refresh_from_db()
            self.assertEqual(item.state, "failed")
            earnings.retry(item.pk)
        claimed = earnings.claim()
        self.assertEqual(claimed.pk, item.pk)
        self.assertIsNone(earnings.claim())
        EarningsDelivery.objects.filter(pk=item.pk).update(started_at=timezone.now()-timedelta(minutes=6))
        self.assertIsNone(earnings.claim())
        item.refresh_from_db()
        self.assertEqual(item.state, "unknown")

    def test_changed_recipient_and_disabled_delivery(self):
        item = earnings.schedule(self.now)
        new = TelegramContact.objects.create(user_id=888, name="Новый")
        Settings.objects.filter(pk=1).update(approver=new)
        self.assertIsNone(earnings.claim())
        item.refresh_from_db()
        self.assertEqual(item.state, "failed")
        self.assertEqual(earnings.retry(item.pk).recipient, 888)
        Settings.objects.filter(pk=1).update(earnings_enabled=False)
        self.assertIsNone(earnings.claim())
        self.assertEqual(EarningsDelivery.objects.get(pk=item.pk).state, "pending")

    def test_worker_runs_earnings_when_invoice_schedule_fails(self):
        from .management.commands.run_billing_worker import Command
        with patch("apps.ledger.management.commands.run_billing_worker.accrue"), \
             patch("apps.ledger.management.commands.run_billing_worker.schedule", side_effect=ValidationError("Invoice error")), \
             patch("apps.ledger.management.commands.run_billing_worker.earnings.schedule") as schedule, \
             self.settings(TELEGRAM_BOT_TOKEN=""), self.assertLogs(level="ERROR"), self.assertRaises(ValidationError):
            async_to_sync(Command().run)(True)
        schedule.assert_called_once()
