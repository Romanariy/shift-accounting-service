from apps.shifts.dates import date_label, date_time_label, excel_date_format
"""Employee earnings from saved ledger amounts; independent of owner invoices."""
import re
from calendar import monthrange
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from asgiref.sync import sync_to_async
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from apps.shifts.audit import log_change
from .billing import safe_text
from .engine import lock
from .models import EarningsDelivery, Record

REPORT_ZONE = ZoneInfo("Asia/Yekaterinburg")


def report_data(month, now=None):
    now = (now or timezone.now()).astimezone(REPORT_ZONE)
    if not isinstance(month, str) or not re.fullmatch(r"\d{4}-\d{2}", month):
        raise ValidationError("Укажите месяц в формате ГГГГ-ММ.")
    try:
        start = date.fromisoformat(month + "-01")
    except ValueError:
        raise ValidationError("Некорректный месяц.")
    if start > now.date():
        raise ValidationError("Будущий месяц недоступен.")
    end = min(start.replace(day=monthrange(start.year, start.month)[1]), now.date())
    employees, review, unassigned = {}, [], []
    rows = Record.objects.filter(date__range=(start, end), deleted_at=None, kind__in=("service", "oneoff")).select_related("employee", "organization").order_by("date", "id")
    for r in rows:
        reasons = []
        if not r.employee_id:
            reasons.append("Не назначен сотрудник")
        if r.review:
            reasons.append("Нужна проверка")
        if r.error:
            reasons.append(r.error)
        if reasons:
            item = {"id": r.pk, "date": r.date.isoformat(), "organization_name": r.organization.name if r.organization_id else "Без организации",
                    "employee_name": r.employee_name or (r.employee.display_name if r.employee_id else "Без сотрудника"),
                    "amount": str(r.amount), "reason": "; ".join(reasons)}
            (unassigned if r.source == "auto" and not r.employee_id else review).append(item)
        else:
            item = employees.setdefault(r.employee_id, {"employee_id": r.employee_id, "employee_name": r.employee.display_name, "amount": Decimal(0)})
            item["amount"] += r.amount
    earnings = sorted(employees.values(), key=lambda r: (r["employee_name"].casefold(), r["employee_id"]))
    total = sum((r["amount"] for r in earnings), Decimal(0))
    for r in earnings:
        r["amount"] = f'{r["amount"]:.2f}'
    return {"month": month, "start": start.isoformat(), "end": end.isoformat(), "generated_at": now.isoformat(),
            "employees": earnings, "total": f"{total:.2f}", "review": review, "unassigned": unassigned,
            "review_total": f'{sum((Decimal(r["amount"]) for r in review), Decimal(0)):.2f}',
            "unassigned_total": f'{sum((Decimal(r["amount"]) for r in unassigned), Decimal(0)):.2f}'}


def workbook_bytes(data):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    main = wb.active
    main.title = "Заработок"
    review = wb.create_sheet("На проверку")
    for ws in (main, review):
        ws.append(["Главный отчёт · все организации"])
        ws.append(["Период", f'{date_label(data["start"])} — {date_label(data["end"])}'])
        ws.append(["Сформирован (Екатеринбург)", date_time_label(data["generated_at"])])
        ws.append([])
    main.append(["Сотрудник", "Заработано, ₽"])
    for r in data["employees"]:
        main.append([safe_text(r["employee_name"]), Decimal(r["amount"])])
    main.append(["Итого заработано", Decimal(data["total"])])
    main.append([])
    main.append(["На проверку (не включено)", Decimal(data["review_total"])])
    main.append(["Начисления без сотрудника (не включено)", Decimal(data["unassigned_total"])])
    review.append(["Дата", "Организация", "Исполнитель", "Сумма, ₽", "Причина", "Категория"])
    for key, label in (("review", "На проверку"), ("unassigned", "Начисления без сотрудника")):
        for r in data[key]:
            review.append([date.fromisoformat(r["date"]), safe_text(r["organization_name"]), safe_text(r["employee_name"]), Decimal(r["amount"]), safe_text(r["reason"]), label])
        review.append([label + " — итого", None, None, Decimal(data[key + "_total"])])
    for ws, widths in ((main, [48, 48]), (review, [22, 30, 30, 20, 55, 36])):
        ws.freeze_panes = "A6"
        ws.sheet_view.showGridLines = False
        for cell in ws[5]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="28756A")
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        for row in ws:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, date):
                    cell.number_format = excel_date_format(cell.value)
                elif isinstance(cell.value, Decimal):
                    cell.number_format = '#,##0.00'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    result = BytesIO()
    wb.save(result)
    return result.getvalue()


def delivery_data(item):
    if not item:
        return None
    return {"id": item.pk, "month": item.month, "recipient": item.recipient, "state": item.state,
            "attempts": item.attempts, "error": item.error, "message_id": item.message_id,
            "created_at": item.created_at.isoformat(), "updated_at": item.updated_at.isoformat()}


@transaction.atomic
def schedule(now=None):
    config = lock()
    now = (now or timezone.now()).astimezone(REPORT_ZONE)
    if not config.earnings_enabled or not config.approver_id or (now.day == 1 and now.hour < 10):
        return None
    month = (now.date().replace(day=1) - timedelta(days=1)).strftime("%Y-%m")
    existing = EarningsDelivery.objects.filter(month=month).first()
    if existing:
        return existing
    data = report_data(month, now=now)
    item = EarningsDelivery.objects.create(month=month, snapshot=data, artifact=workbook_bytes(data), recipient=config.approver.user_id)
    log_change("earnings_delivery", item.pk, "create", actor="system", after={"month": month, "recipient": item.recipient})
    return item


@transaction.atomic
def claim():
    config = lock()
    EarningsDelivery.objects.filter(state="sending", started_at__lt=timezone.now() - timedelta(minutes=5)).update(
        state="unknown", error="Процесс прервался во время доставки. Проверьте Telegram перед повтором.", updated_at=timezone.now())
    if not config.earnings_enabled or not config.approver_id:
        return None
    for item in EarningsDelivery.objects.filter(state="pending").order_by("id"):
        if item.recipient != config.approver.user_id:
            item.state, item.error = "failed", "Утверждающий изменился. Повторите отправку текущему получателю на сайте."
            item.save()
            continue
        item.state, item.started_at, item.attempts = "sending", timezone.now(), item.attempts + 1
        item.save()
        return item
    return None


def finish(pk, state, error="", message_id=None):
    EarningsDelivery.objects.filter(pk=pk, state="sending").update(state=state, error=error[:2000], message_id=message_id, updated_at=timezone.now())


@transaction.atomic
def retry(pk, confirm_duplicate_risk=False):
    config = lock()
    item = EarningsDelivery.objects.get(pk=pk)
    if not config.earnings_enabled or not config.approver_id:
        raise ValidationError("Включите рассылку главного отчёта и назначьте утверждающего.")
    if item.state not in ("failed", "unknown"):
        raise ValidationError("Повтор возможен только для неуспешной или неопределённой доставки.")
    if item.state == "unknown" and confirm_duplicate_risk is not True:
        raise ValidationError("Доставка могла состояться. Подтвердите риск повторного сообщения.")
    item.state, item.error, item.recipient = "pending", "", config.approver.user_id
    item.save()
    log_change("earnings_delivery", item.pk, "retry", actor="web", after={"recipient": item.recipient})
    return item


async def deliver_once(bot):
    from aiogram.types import BufferedInputFile
    from aiogram.exceptions import TelegramBadRequest, TelegramForbiddenError, TelegramRetryAfter, TelegramUnauthorizedError
    item = await sync_to_async(claim)()
    if not item:
        return False
    try:
        data = item.snapshot
        caption = f'Главный отчёт · {item.month}\nВсе организации\nЗаработано сотрудниками: {Decimal(data["total"]):,.2f} ₽\nНа проверку: {len(data["review"])} записей\nНачисления без сотрудника: {len(data["unassigned"])} записей'
        result = await bot.send_document(item.recipient, BufferedInputFile(bytes(item.artifact), filename=f"earnings-{item.month}.xlsx"), caption=caption)
        await sync_to_async(finish)(item.pk, "sent", message_id=result.message_id)
    except TelegramRetryAfter as error:
        await sync_to_async(finish)(item.pk, "failed", f"Лимит Telegram. Повторите через {error.retry_after} секунд.")
    except (TelegramBadRequest, TelegramForbiddenError, TelegramUnauthorizedError) as error:
        await sync_to_async(finish)(item.pk, "failed", str(error))
    except Exception:
        await sync_to_async(finish)(item.pk, "unknown", "Telegram не подтвердил результат доставки. Проверьте чат перед повтором.")
    return True
