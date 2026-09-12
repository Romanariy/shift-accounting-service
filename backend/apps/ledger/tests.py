import json
from datetime import date, datetime, timedelta, timezone as dt_timezone
from decimal import Decimal
from io import BytesIO
from unittest.mock import AsyncMock
from types import SimpleNamespace

from asgiref.sync import async_to_sync
from django.core.exceptions import ValidationError
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from apps.shifts.models import Employee, Organization, ShiftEntry, CompanionEntry, SyncOutbox
from .billing import approve, prepare, refresh, schedule
from .engine import accrue, recalculate, save_record
from .ingest import ingest, parse, source_kind
from .models import Accrual, Batch, Delivery, Invoice, OrganizationBilling, Rate, Record, Service, Settings, TelegramContact
from .worker import claim, deliver_once


class LedgerTests(TestCase):
    def setUp(self):
        self.focus = Organization.objects.get(name="Фокус")
        self.other = Organization.objects.exclude(pk=self.focus.pk).first()
        self.employee = Employee.objects.first()
        self.employee.telegram_user_id = 1234
        self.employee.save()
        self.day = date(2026, 4, 12)
        self.service = Service.objects.get(legacy_code="small_admin")
        self.employee.default_work_type = "small_admin"
        self.employee.save()
        self.companion = Service.objects.get(legacy_code="companion")
        self.companion.default_organization = self.other
        self.companion.save()
        self.recipient = TelegramContact.objects.create(user_id=111, name="Владелец")
        self.approver = TelegramContact.objects.create(user_id=222, name="Утверждающий")
        for org in (self.focus, self.other):
            profile = OrganizationBilling.objects.get(organization=org)
            profile.recipients.add(self.recipient)
        Settings.objects.filter(pk=1).update(approver=self.approver)

    def ingest(self, text, **kwargs):
        return ingest(text, chat_id=-100, message_id=kwargs.pop("message_id", 1), user_id=1234,
                      username="author", author_name="Автор", today=self.day, **kwargs)

    def oneoff(self, amount="1500", **kwargs):
        return self.ingest(f"12.04, {amount}, Починил дверь, Фокус", **kwargs)[0]

    def batch(self):
        return prepare(self.day, self.day, [self.focus.pk])

    def test_split_org_defaults_and_override(self):
        rows = self.ingest("12.04 10:00–14:00 Фокус + 2 сопр")
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0].organization_id, self.focus.pk)
        self.assertEqual(rows[0].units, Decimal(4))
        self.assertEqual(rows[1].organization_id, self.other.pk)
        self.assertEqual(rows[1].date, self.day)
        self.assertEqual(self.ingest("2 сопр Фокус", message_id=2)[0].organization_id, self.focus.pk)
        self.assertEqual(ShiftEntry.objects.filter(telegram_message_id=1).count(), 1)
        self.assertEqual(CompanionEntry.objects.filter(telegram_message_id=1).count(), 1)
        self.assertEqual(SyncOutbox.objects.count(), 3)

    def test_atomic_error_and_no_inherited_org(self):
        self.companion.default_organization = None
        self.companion.save()
        with self.assertRaisesMessage(ValidationError, "Часть 2"):
            self.ingest("12.04 10:00–14:00 Фокус + 2 сопр")
        self.assertEqual(Record.objects.count(), 0)

    def test_oneoff_author_explicit_unknown_and_punctuation(self):
        count = Service.objects.count()
        row = self.ingest("12.04, 1500,50, Починил дверь, замок + ручку, Фокус")[0]
        self.assertEqual(row.amount, Decimal("1500.50"))
        self.assertEqual(row.employee_id, self.employee.pk)
        self.assertEqual(row.description, "Починил дверь, замок + ручку")
        self.assertIsNone(row.service_id)
        self.assertEqual(Service.objects.count(), count)
        unknown = self.ingest("Новый человек, 12.04, 10, Работа, Фокус", message_id=2)[0]
        self.assertTrue(unknown.review)
        self.assertEqual(unknown.employee_name, "Новый человек")
        named = self.ingest(f"{self.employee.short_name}, 12.04, 100, Работа, Фокус", message_id=3)[0]
        self.assertFalse(named.review)

    def test_invalid_oneoff_not_saved(self):
        for text in ("12.04, -20, Работа, Фокус", "31.02, 20, Работа, Фокус", "12.04, 0, Работа, Фокус", "12.04, 20, Работа, Неизвестно"):
            with self.subTest(text=text), self.assertRaises(ValidationError):
                self.ingest(text)
        self.assertFalse(Record.objects.exists())

    def test_expenses_decimal_and_separate_routing(self):
        row = self.ingest("12.04, Краска, кисть, 1500,50, Фокус", expense=True)[0]
        self.assertEqual(row.kind, "expense")
        self.assertEqual(row.description, "Краска, кисть")
        self.assertEqual(row.amount, Decimal("1500.50"))
        self.assertIsNone(row.employee_id)
        Settings.objects.filter(pk=1).update(service_chat=-100, service_thread=1, expense_chat=-100, expense_thread=2)
        self.assertEqual(source_kind(-100, 2), "expense")
        self.assertIsNone(source_kind(-100, 3))

    def test_dedupe_edit_and_frozen_edit(self):
        r = self.oneoff()
        self.oneoff()
        self.assertEqual(Record.objects.count(), 1)
        updated = self.oneoff("2200", edited=True)
        self.assertEqual(updated.pk, r.pk)
        self.assertEqual(updated.amount, Decimal(2200))
        b = self.batch()
        approve(b.pk, b.version)
        with self.assertRaises(ValidationError):
            self.oneoff("3000", edited=True)
        # A retry of the original update is still idempotent even after approval.
        self.oneoff("2200")
        self.assertEqual(Record.objects.count(), 1)

    def test_shorter_edited_message_soft_deletes_extra_parts(self):
        self.ingest("10:00–14:00 Фокус + 2 сопр")
        self.ingest("10:00–15:00 Фокус", edited=True)
        self.assertEqual(Record.objects.filter(deleted_at=None).count(), 1)
        self.assertEqual(CompanionEntry.objects.filter(deleted_at=None).count(), 0)

    def test_rate_overlap_validation_and_night_shift(self):
        rate = Rate.objects.filter(service=self.service, organization=self.focus).first()
        overlapping = Rate(service=self.service, organization=self.focus, price=100, start=self.day)
        with self.assertRaises(ValidationError):
            overlapping.full_clean()
        rate.calculation, rate.price, rate.minimum, rate.maximum = "hourly", Decimal(100), Decimal(600), Decimal(1200)
        rate.save()
        row = self.ingest("12.04 22:00–02:00 Фокус")[0]
        self.assertEqual(row.units, Decimal(4))
        self.assertEqual(row.amount, Decimal(600))

    def test_auto_daily_fixed_idempotence_and_missing_rate(self):
        phone = Service.objects.get(legacy_code="phone")
        Rate.objects.create(service=phone, organization=self.focus, calculation="fixed", price=50, start=self.day)
        Accrual.objects.create(service=phone, organization=self.focus, start=self.day)
        accrue(self.day + timedelta(days=2))
        accrue(self.day + timedelta(days=2))
        rows = Record.objects.filter(service=phone)
        self.assertEqual(rows.count(), 3)
        self.assertEqual(sum(r.amount for r in rows), Decimal(150))
        self.assertFalse(rows.filter(review=True).exists())
        Accrual.objects.create(service=phone, organization=self.other, start=self.day)
        accrue(self.day)
        missing = Record.objects.get(service=phone, organization=self.other)
        self.assertTrue(missing.error)

    def test_xlsx_main_custom_expenses_and_formula_injection(self):
        self.oneoff()
        self.service.sheet = "Смены"
        self.service.save()
        self.ingest("10:00–14:00 Фокус", message_id=2)
        self.ingest("=HYPERLINK(опасно), 20, Фокус", expense=True, message_id=3)
        batch = self.batch()
        invoice = batch.invoices.get()
        workbook = load_workbook(BytesIO(bytes(invoice.artifact)), data_only=False)
        self.assertEqual(workbook.sheetnames, ["Основной", "Смены", "Расходы"])
        main_values = [cell.value for row in workbook["Основной"] for cell in row]
        self.assertIn("Починил дверь", main_values)
        self.assertNotIn("Малый админ", main_values)
        self.assertFalse(any(c.data_type == "f" for ws in workbook for row in ws for c in row))
        self.assertEqual(invoice.total, sum(Decimal(r["amount"]) for r in invoice.lines))
        self.assertEqual(workbook["Основной"]["B3"].value, float(invoice.total))

    def test_refresh_before_approval_freeze_repeat_and_auth(self):
        self.oneoff()
        batch = self.batch()
        with self.assertRaises(ValidationError):
            approve(batch.pk, batch.version, telegram_user=333)
        with self.assertRaises(ValidationError):
            approve(batch.pk, batch.version - 1)
        self.oneoff("500", message_id=2)
        result = approve(batch.pk, batch.version, telegram_user=222)
        self.assertTrue(result["refreshed"])
        self.assertFalse(Delivery.objects.filter(purpose="owner").exists())
        batch.refresh_from_db()
        self.assertTrue(approve(batch.pk, batch.version, telegram_user=222)["approved"])
        frozen = bytes(batch.invoices.get().artifact)
        approve(batch.pk, batch.version)
        self.assertEqual(Delivery.objects.filter(purpose="owner").count(), 1)
        self.service.sheet = "Изменено"
        self.service.save()
        self.assertEqual(bytes(batch.invoices.get().artifact), frozen)
        with self.assertRaises(ValidationError):
            refresh(batch)

    def test_unreviewed_blocks_and_replacement(self):
        self.ingest("Неизвестный, 12.04, 100, Работа, Фокус")
        batch = self.batch()
        with self.assertRaises(ValidationError):
            approve(batch.pk, batch.version)
        row = Record.objects.get()
        save_record({"employee": self.employee.pk}, row)
        refresh(batch)
        approve(batch.pk, batch.version)
        with self.assertRaises(ValidationError):
            self.batch()
        save_record({"amount": "200"}, row, allow_frozen=True)
        replacement = prepare(self.day, self.day, [self.focus.pk], replaces=batch.pk)
        self.assertEqual(replacement.invoices.get().replaces_id, batch.invoices.get().pk)
        approve(replacement.pk, replacement.version)
        self.assertEqual(batch.invoices.get().total, 100)
        self.assertEqual(replacement.invoices.get().total, 200)

    def test_recalc_preview_and_stale_fingerprint(self):
        row = self.ingest("10:00–14:00 Фокус")[0]
        Rate.objects.filter(service=self.service, organization=self.focus).update(price=1000, maximum=None)
        preview = recalculate(self.day, self.day)
        self.assertEqual(len(preview["changes"]), 1)
        row.refresh_from_db()
        self.assertNotEqual(row.amount, 4000)
        Rate.objects.filter(service=self.service, organization=self.focus).update(price=2000)
        with self.assertRaises(ValidationError):
            recalculate(self.day, self.day, preview["fingerprint"])

    def test_recalc_repairs_phone_accruals_with_archived_global_history(self):
        phone = Service.objects.get(legacy_code="phone")
        tomorrow = self.day + timedelta(days=1)
        Rate.objects.filter(service=phone).delete()
        for org in (self.focus, self.other):
            Accrual.objects.create(service=phone, organization=org, start=self.day)
        accrue(tomorrow)
        rows = Record.objects.filter(service=phone, source="auto")
        self.assertEqual(rows.count(), 4)
        self.assertEqual(rows.filter(amount=0, review=True).count(), 4)
        archive = Record.objects.create(kind="service", service=phone, date=self.day, amount=200,
            source="import", review=True, description="Архив общего расчёта телефонов",
            auto_key=f"legacy-phone:{self.day}")
        Rate.objects.create(service=phone, organization=self.focus, calculation="fixed", price=100,
                            start=self.day, end=self.day)
        Rate.objects.create(service=phone, organization=self.focus, calculation="fixed", price=250,
                            start=tomorrow)
        Rate.objects.create(service=phone, organization=self.other, calculation="fixed", price=300,
                            start=self.day)
        endpoint = "/api/shifts/ledger/recalculate/"
        payload = {"start": str(self.day), "end": str(tomorrow)}
        response = self.client.post(endpoint, json.dumps(payload), content_type="application/json")
        self.assertEqual(response.status_code, 200, response.content)
        preview = response.json()
        self.assertEqual(preview["errors"], [])
        self.assertEqual(len(preview["changes"]), 4)
        self.assertEqual(preview["skipped"][0]["reason"], "Архив общего расчёта телефонов")
        self.assertEqual(rows.filter(amount=0).count(), 4)  # Preview does not write.
        payload["fingerprint"] = preview["fingerprint"]
        response = self.client.post(endpoint, json.dumps(payload), content_type="application/json")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(rows.get(organization=self.focus, date=self.day).amount, 100)
        self.assertEqual(rows.get(organization=self.focus, date=tomorrow).amount, 250)
        self.assertEqual(list(rows.filter(organization=self.other).values_list("amount", flat=True)), [300, 300])
        self.assertFalse(rows.filter(review=True).exists())
        self.assertFalse(rows.exclude(error="").exists())
        self.assertEqual(rows.count(), 4)
        archive.refresh_from_db()
        self.assertEqual(archive.amount, 200)
        self.assertIsNone(archive.organization_id)
        self.assertTrue(archive.review)
        self.assertFalse(Delivery.objects.exists())
        self.assertEqual(recalculate(self.day, tomorrow)["changes"], [])

    def test_recalc_unassigned_history_does_not_hide_missing_current_tariff(self):
        phone = Service.objects.get(legacy_code="phone")
        archive = Record.objects.create(kind="service", service=phone, date=self.day, amount=200)
        Rate.objects.filter(service=phone).delete()
        Accrual.objects.create(service=phone, organization=self.focus, start=self.day)
        accrue(self.day)
        preview = recalculate(self.day, self.day)
        self.assertEqual(preview["skipped"][0]["reason"], "Не назначена организация")
        self.assertEqual(len(preview["errors"]), 1)
        with self.assertRaisesMessage(ValidationError, "Сначала исправьте отсутствующие тарифы"):
            recalculate(self.day, self.day, preview["fingerprint"])
        archive.refresh_from_db()
        self.assertEqual(archive.amount, 200)

    def test_api_bootstrap_crud_and_summary(self):
        self.assertEqual(self.client.get("/api/shifts/ledger/bootstrap/").status_code, 200)
        payload = {"kind": "oneoff", "organization": self.focus.pk, "employee": self.employee.pk,
                   "date": self.day.isoformat(), "amount": "100", "description": "Разовая"}
        response = self.client.post("/api/shifts/ledger/records/", json.dumps(payload), content_type="application/json")
        self.assertEqual(response.status_code, 200, response.content)
        data = self.client.get("/api/shifts/ledger/records/").json()
        self.assertEqual(data["employee_totals"][0]["total"], "100")
        response = self.client.post("/api/shifts/ledger/services/", json.dumps({"name":"Тест", "aliases":["сопр"]}), content_type="application/json")
        self.assertEqual(response.status_code, 400)

    def test_scheduled_batch_once_and_no_owner_without_approval(self):
        Settings.objects.filter(pk=1).update(enabled=True)
        OrganizationBilling.objects.filter(organization=self.focus).update(monthly=True)
        self.oneoff()
        now = datetime(2026,5,1,6,tzinfo=dt_timezone.utc)
        first = schedule(now)
        self.assertEqual(schedule(now).pk, first.pk)
        self.assertEqual(Batch.objects.count(), 1)
        self.assertFalse(Delivery.objects.filter(purpose="owner").exists())

    def test_delivery_unknown_failure_and_partial_progress(self):
        self.oneoff()
        batch = self.batch()
        approve(batch.pk, batch.version)
        bot = AsyncMock()
        bot.send_document.side_effect = TimeoutError()
        async_to_sync(deliver_once)(bot)
        self.assertEqual(Delivery.objects.get(purpose="owner").state, "unknown")
        self.assertFalse(Delivery.objects.filter(purpose="owner", state="pending").exists())
        delivery = Delivery.objects.get(purpose="owner")
        response = self.client.post(f"/api/shifts/ledger/deliveries/{delivery.pk}/retry/", "{}", content_type="application/json")
        self.assertEqual(response.status_code, 400)
        response = self.client.post(f"/api/shifts/ledger/deliveries/{delivery.pk}/retry/", '{"confirm_duplicate_risk": true}', content_type="application/json")
        self.assertEqual(response.status_code, 200)

    def test_recover_crash_does_not_resend(self):
        self.oneoff()
        batch = self.batch()
        approve(batch.pk, batch.version)
        item = claim()
        Delivery.objects.filter(pk=item.pk).update(started_at=timezone.now()-timedelta(minutes=6))
        self.assertIsNone(claim())
        item.refresh_from_db()
        self.assertEqual(item.state, "unknown")

    def test_successful_delivery_and_no_duplicate_after_retry(self):
        self.oneoff()
        batch = self.batch()
        # An approver gets each workbook before the approval buttons.
        bot = AsyncMock()
        bot.send_document.return_value = SimpleNamespace(message_id=901)
        bot.send_message.return_value = SimpleNamespace(message_id=902)
        self.assertTrue(async_to_sync(deliver_once)(bot))
        self.assertEqual(Delivery.objects.get(purpose="preview").state, "sent")
        self.assertTrue(async_to_sync(deliver_once)(bot))
        self.assertEqual(Delivery.objects.get(purpose="approval").state, "sent")
        approve(batch.pk, batch.version)
        self.assertTrue(async_to_sync(deliver_once)(bot))
        self.assertFalse(async_to_sync(deliver_once)(bot))
        self.assertEqual(Delivery.objects.get(purpose="owner").message_id, 901)
        self.assertEqual(bot.send_document.call_count, 2)

    def test_partial_delivery_only_failed_recipient_retried(self):
        self.oneoff()
        second = TelegramContact.objects.create(user_id=444, name="Второй владелец")
        OrganizationBilling.objects.get(organization=self.focus).recipients.add(second)
        batch = self.batch()
        approve(batch.pk, batch.version)
        bot = AsyncMock()
        bot.send_document.side_effect = [SimpleNamespace(message_id=10), TimeoutError()]
        async_to_sync(deliver_once)(bot)
        async_to_sync(deliver_once)(bot)
        self.assertEqual(Delivery.objects.filter(purpose="owner", state="sent").count(), 1)
        self.assertEqual(Delivery.objects.filter(purpose="owner", state="unknown").count(), 1)
        self.assertIsNone(claim())

    def test_main_sheet_services_oneoffs_adjustments_and_isolation(self):
        self.oneoff()
        self.ingest("10:00–14:00 Фокус + 2 сопр", message_id=2)
        batch = self.batch()
        invoice = batch.invoices.get()
        response = self.client.put(f"/api/shifts/ledger/invoices/{invoice.pk}/", json.dumps({"adjustments":[{"description":"Скидка", "amount":"-50"}]}), content_type="application/json")
        self.assertEqual(response.status_code, 200, response.content)
        invoice.refresh_from_db()
        wb = load_workbook(BytesIO(bytes(invoice.artifact)))
        self.assertEqual(wb.sheetnames, ["Основной"])
        values = [c.value for row in wb.active for c in row]
        self.assertIn("Починил дверь", values)
        self.assertIn("Малый админ", values)
        self.assertNotIn("Сопровождение", values)
        self.assertIn("Скидка", values)
        self.assertEqual(values.count("Починил дверь"), 1)

    def test_two_drafts_cannot_both_bill_same_period(self):
        self.oneoff()
        first, second = self.batch(), self.batch()
        approve(first.pk, first.version)
        with self.assertRaises(ValidationError):
            approve(second.pk, second.version)

    def test_monthly_accrual_and_no_employee_required(self):
        service = Service.objects.create(name="Абонемент", input_type="mark", frequency="monthly")
        Rate.objects.create(service=service, organization=self.focus, calculation="fixed", price=500, start=self.day)
        Accrual.objects.create(service=service, organization=self.focus, start=self.day)
        accrue(date(2026, 6, 10))
        rows = Record.objects.filter(service=service)
        self.assertEqual(rows.count(), 3)
        self.assertFalse(rows.filter(review=True).exists())
        self.assertEqual(sum(r.amount for r in rows), 1500)

    def test_resolving_review_and_recalc_preserves_frozen(self):
        row = self.ingest("10:00–14:00 Фокус")[0]
        batch = self.batch()
        approve(batch.pk, batch.version)
        Rate.objects.filter(service=self.service, organization=self.focus).update(price=999)
        self.assertEqual(recalculate(self.day, self.day)["changes"], [])
        legacy = ShiftEntry.objects.get(pk=row.legacy_id)
        legacy.calculated_amount = 12345
        with self.assertRaises(ValidationError):
            legacy.save()

    def test_legacy_employee_summary_includes_oneoff_not_expense(self):
        self.oneoff()
        self.ingest("Краска, 200, Фокус", message_id=2, expense=True)
        response = self.client.get("/api/shifts/month-summary/?year=2026&month=4")
        self.assertEqual(response.status_code, 200, response.content)
        employee = next(e for e in response.json()["employees"] if e["employeeId"] == self.employee.pk)
        self.assertEqual(Decimal(employee["additionalAmount"]), 1500)

    def test_missing_tariff_blocks_existing_record_invoice(self):
        self.ingest("10:00–14:00 Фокус")
        Rate.objects.filter(service=self.service, organization=self.focus).update(active=False)
        batch = self.batch()
        self.assertTrue(batch.invoices.get().errors)
        with self.assertRaises(ValidationError):
            approve(batch.pk, batch.version)

    def test_history_includes_original_legacy_audit(self):
        from apps.shifts.audit import log_change
        row = self.ingest("10:00–14:00 Фокус")[0]
        original = log_change("shift", row.legacy_id, "create", actor="old-import", after={"amount":"123"})
        response = self.client.get(f"/api/shifts/ledger/audit/?entity_type=record&entity_id={row.pk}")
        self.assertEqual(response.status_code, 200)
        self.assertIn(original.pk, [item["id"] for item in response.json()])
        item = next(item for item in response.json() if item["id"] == original.pk)
        self.assertEqual(datetime.fromisoformat(item["created_at"]), original.created_at)
        self.assertEqual(item["actor"], "old-import")
        self.assertEqual(item["diff"]["amount"], {"from": None, "to": "123"})

    def test_history_is_isolated_and_sorted_with_tied_timestamps(self):
        from apps.shifts.audit import log_change
        from apps.shifts.models import AuditLog
        row = self.oneoff()
        other = self.oneoff(message_id=2)
        first = log_change("record", row.pk, "update", before={"amount": "1500"}, after={"amount": "1600"})
        last = log_change("record", row.pk, "update", before={"amount": "1600"}, after={"amount": "1700"})
        unrelated = log_change("record", other.pk, "update", after={"amount": "9999"})
        AuditLog.objects.filter(pk__in=[first.pk, last.pk]).update(created_at=timezone.now())
        response = self.client.get("/api/shifts/ledger/audit/", {"entity_type": "record", "entity_id": row.pk})
        self.assertEqual(response.status_code, 200)
        ids = [item["id"] for item in response.json()]
        self.assertEqual(ids[:2], [last.pk, first.pk])
        self.assertNotIn(unrelated.pk, ids)
        empty = self.client.get("/api/shifts/ledger/audit/", {"entity_type": "record", "entity_id": 999999})
        self.assertEqual(empty.json(), [])

    def test_month_filters_search_cyrillic_and_sort_before_pagination(self):
        low = self.oneoff("100")
        high = self.oneoff("900", message_id=2)
        middle = self.oneoff("500", message_id=3)
        other_month = self.oneoff("9999", message_id=4)
        Record.objects.filter(pk=other_month.pk).update(date=date(2026, 3, 31))
        response = self.client.get("/api/shifts/ledger/records/", {"month":"2026-04", "q":"ПОЧИНИЛ", "ordering":"-amount", "limit":1, "offset":1})
        self.assertEqual(response.status_code, 200, response.content)
        result = response.json()
        self.assertEqual(result["count"], 3)
        self.assertEqual(Decimal(result["total"]), 1500)
        self.assertEqual([r["id"] for r in result["records"]], [middle.pk])
        self.assertEqual(result["limit"], 1)
        result = self.client.get("/api/shifts/ledger/records/", {"month":"2026-04", "ordering":"amount"}).json()
        self.assertEqual([r["id"] for r in result["records"]], [low.pk, middle.pk, high.pk])

    def test_journal_filters_combine_and_validate_parameters(self):
        row = self.oneoff()
        self.ingest("Незнакомый, 12.04, 300, Починил замок, Фокус", message_id=2)
        self.ingest("Краска, 200, Фокус", message_id=3, expense=True)
        params = {"month":"2026-04", "organization":self.focus.pk, "employee":self.employee.pk, "kind":"oneoff", "status":"ready", "source":"telegram"}
        data = self.client.get("/api/shifts/ledger/records/", params).json()
        self.assertEqual(data["count"], 1)
        self.assertEqual(data["records"][0]["id"], row.pk)
        data = self.client.get("/api/shifts/ledger/records/", {"month":"2026-04", "employee":"unassigned", "kind":"work", "status":"review"}).json()
        self.assertEqual(data["count"], 1)
        self.assertEqual(data["review"], 1)
        for invalid in ({"month":"2026-99"}, {"month":"2026-4"}, {"ordering":"raw_text"}, {"status":"fake"}):
            self.assertEqual(self.client.get("/api/shifts/ledger/records/", invalid).status_code, 400)

    def test_all_records_automatically_count_despite_old_exclusion_flags(self):
        work = self.oneoff()
        expense = self.ingest("Краска, 200, Фокус", expense=True, message_id=2)[0]
        shift = self.ingest("10:00–14:00 Фокус", message_id=3)[0]
        Record.objects.filter(pk__in=[work.pk, expense.pk, shift.pk]).update(included=False)
        self.service.included = False
        self.service.save()
        profile = OrganizationBilling.objects.get(organization=self.focus)
        profile.include_expenses = False
        profile.save()
        profile.excluded_services.add(self.service)
        batch = self.batch()
        invoice = batch.invoices.get()
        self.assertEqual({line["id"] for line in invoice.lines}, {work.pk, expense.pk, shift.pk})
        self.assertEqual(invoice.total, work.amount + expense.amount + shift.amount)
        self.assertEqual(invoice.excluded_ids, [])
        response = self.client.put(f"/api/shifts/ledger/invoices/{invoice.pk}/", json.dumps({"excluded_ids":[work.pk]}), content_type="application/json")
        self.assertEqual(response.status_code, 400)

    def test_delete_approved_invoice_removes_document_queue_and_preserves_work(self):
        row = self.oneoff()
        batch = self.batch()
        approve(batch.pk, batch.version)
        invoice = batch.invoices.get()
        response = self.client.delete(f"/api/shifts/ledger/invoices/{invoice.pk}/")
        self.assertEqual(response.status_code, 200, response.content)
        self.assertFalse(Invoice.objects.filter(pk=invoice.pk).exists())
        self.assertFalse(Batch.objects.filter(pk=batch.pk).exists())
        self.assertFalse(Delivery.objects.filter(batch_id=batch.pk).exists())
        self.assertTrue(Record.objects.filter(pk=row.pk, deleted_at=None).exists())
        self.assertEqual(self.client.get(f"/api/shifts/ledger/invoices/{invoice.pk}/file/").status_code, 404)
        save_record({"amount":"200"}, row)
        self.assertEqual(row.amount, 200)

    def test_delete_one_invoice_keeps_other_organization_and_invalidates_draft(self):
        self.oneoff()
        batch = prepare(self.day, self.day, [self.focus.pk, self.other.pk])
        previous_version = batch.version
        invoice = batch.invoices.get(organization=self.focus)
        response = self.client.delete(f"/api/shifts/ledger/invoices/{invoice.pk}/")
        self.assertEqual(response.status_code, 200, response.content)
        batch.refresh_from_db()
        self.assertEqual(batch.organization_ids, [self.other.pk])
        self.assertEqual(batch.invoices.count(), 1)
        self.assertGreater(batch.version, previous_version)
        with self.assertRaises(ValidationError):
            approve(batch.pk, previous_version)

    def test_delete_blocks_inflight_delivery_and_allows_sent_document_removal(self):
        self.oneoff()
        batch = self.batch()
        approve(batch.pk, batch.version)
        delivery = claim()
        self.assertEqual(self.client.delete(f"/api/shifts/ledger/batches/{batch.pk}/").status_code, 400)
        Delivery.objects.filter(pk=delivery.pk).update(state="sent", message_id=123)
        self.assertEqual(self.client.delete(f"/api/shifts/ledger/batches/{batch.pk}/").status_code, 200)
        self.assertFalse(Invoice.objects.filter(batch_id=batch.pk).exists())

    def test_deleted_scheduled_batch_is_not_automatically_recreated(self):
        Settings.objects.filter(pk=1).update(enabled=True)
        OrganizationBilling.objects.filter(organization=self.focus).update(monthly=True)
        self.oneoff()
        now = datetime(2026,5,1,6,tzinfo=dt_timezone.utc)
        batch = schedule(now)
        self.assertEqual(self.client.delete(f"/api/shifts/ledger/batches/{batch.pk}/").status_code, 200)
        self.assertIsNone(schedule(now))
        self.assertFalse(Batch.objects.exists())

    def test_delete_replaced_invoice_preserves_newer_version(self):
        self.oneoff()
        original = self.batch()
        approve(original.pk, original.version)
        replacement = prepare(self.day, self.day, [self.focus.pk], replaces=original.pk)
        self.assertEqual(self.client.delete(f"/api/shifts/ledger/batches/{original.pk}/").status_code, 200)
        self.assertIsNone(replacement.invoices.get().replaces_id)
        self.assertTrue(approve(replacement.pk, replacement.version)["approved"])

    def test_invoice_list_month_organization_status_and_sort(self):
        self.oneoff()
        april = self.batch()
        march = prepare(date(2026,3,1), date(2026,3,31), [self.other.pk])
        result = self.client.get("/api/shifts/ledger/batches/", {"month":"2026-04", "organization":self.focus.pk, "state":"draft"})
        self.assertEqual(result.status_code, 200, result.content)
        self.assertEqual([b["id"] for b in result.json()], [april.pk])
        self.assertEqual(self.client.get("/api/shifts/ledger/batches/", {"state":"approved"}).json(), [])
        ids = [b["id"] for b in self.client.get("/api/shifts/ledger/batches/", {"ordering":"id"}).json()]
        self.assertEqual(ids, [april.pk, march.pk])
