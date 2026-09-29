import json
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from openpyxl import load_workbook

from apps.shifts.models import Employee, Organization, AuditLog
from .engine import save_record, preview_record, delete_record, recalculate
from .models import Batch, Invoice, OrganizationReceipt, Rate, Record, Service, OverlapDecision
from .overlaps import review_day
from .billing import prepare, refresh, approve


@override_settings(SHARED_SHIFT_ALLOCATION_ENABLED=False, SHIFT_SYNC_AFTER_WRITE=False, TELEGRAM_BOT_TOKEN="")
class OverlapAndReceiptTests(TestCase):
    def setUp(self):
        self.orgs = [Organization.objects.create(name=f"Студия тест {i}") for i in range(3)]
        self.employee = Employee.objects.create(short_name="Сотрудник пересечений")
        self.service = Service.objects.create(name="Тест смены", legacy_code="small_admin", input_type="time")
        self.day = date(2026, 4, 12)
        for org in self.orgs:
            Rate.objects.create(service=self.service, organization=org, calculation="hourly", price=200,
                                minimum=600, maximum=1500, start=date(2026, 1, 1))

    def shift(self, org, start, end, **extra):
        return save_record({"kind": "service", "service": self.service.pk, "organization": self.orgs[org].pk,
                            "date": str(self.day), "employee": self.employee.pk,
                            "start_time": start, "end_time": end, **extra})

    def post(self, resource, payload):
        return self.client.post(f"/api/shifts/ledger/{resource}/", json.dumps(payload), content_type="application/json")

    def test_one_minute_overlap_halves_clamped_base_and_keeps_bonus(self):
        a = self.shift(0, "08:00", "16:00")
        b = self.shift(1, "15:59", "17:59", bonus="500", bonus_reason="Срочный выход")
        a.refresh_from_db()
        self.assertEqual(a.amount, 1500)
        self.assertEqual((b.base_amount, b.overlap_factor, b.bonus, b.amount), (Decimal(600), Decimal(".5"), Decimal(500), Decimal(800)))
        b = save_record({"description": "Комментарий"}, b)
        self.assertEqual(b.amount, 800)
        save_record({"start_time": "16:00", "end_time": "18:00"}, b)
        b.refresh_from_db()
        self.assertEqual(b.amount, 1100)

    def test_no_overlap_different_employee_and_other_services(self):
        self.shift(0, "08:00", "16:00")
        b = self.shift(1, "16:00", "18:00")
        self.assertEqual(b.overlap_factor, 1)
        other = Employee.objects.create(short_name="Другой")
        b = save_record({"employee": other.pk, "start_time": "12:00", "end_time": "14:00"}, b)
        self.assertEqual(b.amount, 600)
        service = Service.objects.create(name="Уборка тест", input_type="time", legacy_code="cleaning")
        Rate.objects.create(service=service, organization=self.orgs[0], calculation="hourly", price=200, start=self.day)
        c = save_record({"kind":"service", "service":service.pk, "organization":self.orgs[0].pk, "date":str(self.day),
                         "employee":self.employee.pk,"start_time":"10:00","end_time":"12:00"})
        self.assertEqual(c.amount, 400)
        with self.assertRaises(ValidationError):
            save_record({"bonus":"1","bonus_reason":"Не смена"}, c)

    def test_equal_hours_cheaper_base_then_stable_organization(self):
        Rate.objects.filter(service=self.service, organization=self.orgs[1]).update(price=100, minimum=None)
        a = self.shift(0, "10:00", "14:00")
        b = self.shift(1, "10:00", "14:00")
        self.assertEqual(b.amount, 200)
        Rate.objects.filter(service=self.service, organization=self.orgs[1]).update(price=200)
        save_record({"units":"4", "bonus_reason":"Изменение"}, b)
        a.refresh_from_db(); b.refresh_from_db()
        self.assertEqual((a.amount, b.amount), (400, 800))

    def test_all_studio_records_share_factor_delete_restores_other_studio(self):
        a = self.shift(0, "08:00", "12:00")
        a2 = self.shift(0, "14:00", "18:00")
        b = self.shift(1, "11:00", "15:00")
        self.assertEqual(b.amount, 400)
        delete_record(a2)
        a.refresh_from_db(); b.refresh_from_db()
        self.assertEqual((a.amount, b.amount), (400, 800))
        delete_record(a)
        b.refresh_from_db()
        self.assertEqual(b.amount, 800)

    def test_manual_three_studios_and_changed_members_require_new_decision(self):
        rows = [self.shift(i, "10:00", f"{18-i}:00") for i in range(3)]
        for r in rows:
            r.refresh_from_db();self.assertTrue(r.calculation_review)
        component = review_day(self.employee.pk, self.day)[0]
        payload = {**component, "half_organizations": [self.orgs[1].pk,self.orgs[2].pk]}
        response = self.post("overlap-decisions", payload)
        self.assertEqual(response.status_code, 200, response.content)
        self.assertEqual(self.post("overlap-decisions", payload).status_code, 400)
        for r in rows:
            r.refresh_from_db();self.assertFalse(r.calculation_review)
        self.assertEqual(rows[2].amount, 600)
        save_record({"end_time":"15:00"}, rows[2])
        rows[0].refresh_from_db();self.assertTrue(rows[0].calculation_review)
        self.assertEqual(OverlapDecision.objects.count(),1)

    def test_missing_intervals_are_not_guessed(self):
        r = self.shift(0, None, None, units="4")
        self.assertIn("начало", r.calculation_note)
        self.assertEqual(r.overlap_factor, 1)

    def test_preview_rolls_back_and_rejects_changes_to_peer(self):
        a = self.shift(0,"10:00","18:00")
        payload = {"kind":"service","service":self.service.pk,"organization":self.orgs[1].pk,"employee":self.employee.pk,
                   "date":str(self.day),"start_time":"12:00","end_time":"16:00","bonus":"100","bonus_reason":"Выход"}
        count = Record.objects.count()
        preview = preview_record(payload)["record"]
        self.assertEqual(Decimal(preview["amount"]), 500)
        self.assertEqual(Record.objects.count(),count)
        save_record({"end_time":"17:00"},a)
        with self.assertRaisesMessage(ValidationError,"Данные изменились"):
            save_record({**payload,"preview_fingerprint":preview["preview_fingerprint"]})

    def test_frozen_peer_and_old_invoice_artifact_are_preserved(self):
        a = self.shift(0,"10:00","14:00")
        batch = Batch.objects.create(start=self.day,end=self.day,organization_ids=[self.orgs[0].pk],state="approved")
        inv = Invoice.objects.create(batch=batch,organization=self.orgs[0],organization_name=self.orgs[0].name,lines=[{"id":a.pk}],artifact=b"saved")
        with self.assertRaises(ValidationError):
            self.shift(1,"10:00","18:00")
        self.assertEqual(Record.objects.count(),1)
        b = save_record({"kind":"service","service":self.service.pk,"organization":self.orgs[1].pk,"employee":self.employee.pk,
                         "date":str(self.day),"start_time":"10:00","end_time":"18:00"},allow_frozen=True)
        a.refresh_from_db();inv.refresh_from_db()
        self.assertEqual(a.amount,400)
        self.assertEqual(bytes(inv.artifact),b"saved")

    def test_historical_recalculation_preview_and_bonus_idempotence(self):
        a = self.shift(0,"10:00","18:00")
        b = self.shift(1,"12:00","16:00",bonus="100",bonus_reason="Причина")
        Record.objects.filter(pk=b.pk).update(amount=900,base_amount=None,overlap_factor=1)
        preview = recalculate(self.day,self.day)
        b.refresh_from_db();self.assertEqual(b.amount,900)
        recalculate(self.day,self.day,preview["fingerprint"])
        b.refresh_from_db();self.assertEqual(b.amount,500)
        self.assertEqual(recalculate(self.day,self.day)["changes"],[])

    def test_receipts_isolation_history_and_invoice_snapshot(self):
        self.shift(0,"10:00","14:00")
        payload = {"organization":self.orgs[0].pk,"date":str(self.day),"payer":"=Клиент","amount":"1234.50"}
        response=self.post("receipts",payload)
        self.assertEqual(response.status_code,200,response.content)
        receipt=response.json()
        batch=prepare(self.day,self.day,[self.orgs[0].pk]);inv=batch.invoices.get()
        self.assertEqual(inv.total,800)
        wb=load_workbook(BytesIO(bytes(inv.artifact)))
        self.assertEqual(wb["Полученные оплаты"]["B2"].value,"'=Клиент")
        self.assertEqual(wb["Полученные оплаты"]["C3"].value,1234.5)
        self.assertEqual(len(inv.receipts),1)
        batch.state="approved";batch.save()
        original=bytes(inv.artifact)
        response=self.client.put(f'/api/shifts/ledger/receipts/{receipt["id"]}/',json.dumps({**receipt,"amount":"2000"}),content_type="application/json")
        self.assertEqual(response.status_code,200,response.content)
        inv.refresh_from_db();self.assertEqual(bytes(inv.artifact),original)
        self.assertEqual(inv.payment_state,"none")
        listing=self.client.get(f'/api/shifts/ledger/receipts/?organization={self.orgs[1].pk}').json()
        self.assertEqual(listing["count"],0)
        receipt=response.json()
        response=self.client.delete(f'/api/shifts/ledger/receipts/{receipt["id"]}/',json.dumps(receipt),content_type="application/json")
        self.assertEqual(response.status_code,200,response.content)
        self.assertEqual(AuditLog.objects.filter(entity_type="receipt").count(),3)
        self.assertIsNotNone(OrganizationReceipt.objects.get().deleted_at)

    def test_unresolved_calculation_blocks_invoice_approval(self):
        for i in range(3): self.shift(i,"10:00","18:00")
        batch=prepare(self.day,self.day,[self.orgs[0].pk])
        inv=batch.invoices.get()
        self.assertTrue(any("три" in e for e in inv.errors))
        with self.assertRaises(ValidationError):approve(batch.pk,batch.version)

    def test_receipt_rejects_invalid_amount_and_stale_edit(self):
        payload={"organization":self.orgs[0].pk,"date":str(self.day),"payer":"Клиент","amount":"0"}
        self.assertEqual(self.post("receipts",payload).status_code,400)
        r=self.post("receipts",{**payload,"amount":"100"}).json()
        url=f'/api/shifts/ledger/receipts/{r["id"]}/'
        self.assertEqual(self.client.put(url,json.dumps({**r,"amount":"200"}),content_type="application/json").status_code,200)
        self.assertEqual(self.client.put(url,json.dumps(r),content_type="application/json").status_code,400)

    def test_required_comment_hint_and_telegram_edit_keeps_bonus(self):
        from .ingest import ingest
        self.service.comment_required = True
        self.service.comment_hint = "Укажите зал и выполненную работу."
        self.service.save()
        self.employee.telegram_user_id = 7788
        self.employee.save()
        message = "12.04 Тест смены 10:00-14:00 Студия тест 0"
        args = dict(chat_id=-900, message_id=1, user_id=7788, today=self.day)
        with self.assertRaisesMessage(ValidationError, self.service.comment_hint):
            ingest(message, **args)
        self.assertEqual(Record.objects.count(), 0)
        row = ingest(message + " Убрал зал", **args)[0]
        row = save_record({"bonus":"100", "bonus_reason":"Срочная работа"}, row)
        result = ingest(message + " Убрал большой зал", edited=True, **args)[0]
        self.assertEqual(result.bonus, 100)
        self.assertEqual(result.amount, 900)
        self.assertEqual(result.base_amount, 800)
        with self.assertRaisesMessage(ValidationError,self.service.comment_hint):
            save_record({"description":"  "}, result)

    def test_fractional_kopecks_halve_studio_total_once(self):
        Rate.objects.filter(service=self.service).update(minimum=None,maximum=None,price="100.01")
        self.shift(0,"08:00","18:00")
        a=self.shift(1,"10:00","11:00")
        b=self.shift(1,"11:00","12:00")
        a.refresh_from_db()
        self.assertEqual(a.amount+b.amount,Decimal("100.01"))

    def test_deductions_reduce_only_receipt_balance_and_are_snapshotted(self):
        self.shift(0,"10:00","14:00")
        common={"organization":self.orgs[0].pk,"date":str(self.day)}
        self.assertEqual(self.post("receipts",{**common,"payer":"Клиент","amount":"1000"}).status_code,200)
        self.assertEqual(self.post("receipts",{**common,"kind":"deduction","amount":"200"}).status_code,400)
        r=self.post("receipts",{**common,"kind":"deduction","amount":"200","reason":"Возврат клиенту"})
        self.assertEqual(r.status_code,200,r.content)
        result=self.client.get(f'/api/shifts/ledger/receipts/?organization={self.orgs[0].pk}').json()
        self.assertEqual((Decimal(result["received_total"]),Decimal(result["deduction_total"]),Decimal(result["total"])),(1000,200,800))
        inv=prepare(self.day,self.day,[self.orgs[0].pk]).invoices.get()
        self.assertEqual(inv.total,800)
        self.assertEqual(Record.objects.get().amount,800)
        sheet=load_workbook(BytesIO(bytes(inv.artifact)))["Полученные оплаты"]
        self.assertEqual(sheet["C3"].value,-200)
        self.assertEqual(sheet["C6"].value,800)
        self.assertEqual(inv.receipts[1]["reason"],"Возврат клиенту")
