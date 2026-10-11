import json
from datetime import date, time
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch

from django.core.exceptions import ValidationError
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from apps.shifts.models import Organization, Employee, AuditLog
from apps.offers.models import RecognitionJob
from .models import Service, Rate, ClientBooking, ClientBonusCharge, BonusRule, ClientBookingImport
from .engine import save_record, delete_record
from .client_bonuses import (normalize_phone, save_client, save_rule, save_booking, recalculate, data, assessment)
from .booking_imports import sanitize_phones, finish_import, stage_bookings
from .billing import prepare, approve


class ClientBonusTests(TestCase):
    def setUp(self):
        self.day = date(2026, 4, 12)
        self.org = Organization.objects.create(name="Bonus studio")
        self.employee = Employee.objects.create(full_name="Bonus employee", short_name="Bonus employee")
        self.service = Service.objects.create(name="Bonus admin", input_type="time", legacy_code="small_admin")
        Rate.objects.create(service=self.service, organization=self.org, calculation="hourly", price=200, start=self.day)
        self.client1 = save_client({"name": "Школа", "phones": ["8 (919) 123-45-67", "+7 919 123 45 68"]})
        self.client2 = save_client({"name": "Другой клиент", "phones": ["+79195555555"]})
        self.rule = save_rule({"name": "Школьные съёмки", "organization": self.org.pk, "amount": "200", "start": str(self.day),
                              "clients": [self.client1["id"], self.client2["id"]]})
        self.activate()

    def activate(self, rule=None):
        rule = rule or self.rule
        request = {"start": str(self.day), "end": str(self.day), "rule": rule["id"], "revision": rule["revision"]}
        preview = recalculate(request)
        return recalculate({**request, "fingerprint": preview["fingerprint"]})

    def work(self, employee=None, start="10:00", end="18:00", bonus="0"):
        return save_record({"kind": "service", "service": self.service.pk, "organization": self.org.pk,
            "employee": (employee or self.employee).pk, "date": str(self.day), "start_time": start, "end_time": end,
            "bonus": bonus, "bonus_reason": "Подмена" if Decimal(bonus) else ""})

    def booking(self, client=None, start="12:00", end="13:00"):
        obj = save_booking({"client": client or self.client1["id"], "organization": self.org.pk, "date": str(self.day),
                           "start_time": start, "end_time": end, "rooms": ["Зал"]})
        return save_booking({"revision": obj["revision"]}, obj["id"], "confirm")

    def test_multiple_clients_multiple_phones_and_distinct_bookings(self):
        row = self.work(bonus="500")
        self.booking(); self.booking(start="14:00", end="15:00"); self.booking(self.client2["id"], "16:00", "17:00")
        row.refresh_from_db()
        self.assertEqual((row.amount, row.client_bonus, row.bonus), (Decimal("2700"), Decimal("600"), Decimal("500")))
        self.assertEqual(len(row.client_bonus_details), 3)
        duplicate = save_booking({"client": self.client1["id"], "organization": self.org.pk, "date": str(self.day),
                                 "start_time": "12:00", "end_time": "13:00", "rooms": ["Другой зал"]})
        self.assertTrue(duplicate["duplicate"])
        self.assertEqual(ClientBonusCharge.objects.filter(state="applied").count(), 3)
        delete_record(row)
        self.assertFalse(ClientBonusCharge.objects.filter(state="applied").exists())

    def test_normalization_validation_and_revision(self):
        self.assertEqual(normalize_phone("8 (919) 123-45-67"), "+79191234567")
        for value in ("9191234567", "7919123456", "abc+79191234567"):
            with self.assertRaises(ValidationError): normalize_phone(value)
        with self.assertRaises(ValidationError): save_client({"phones": ["+79191234567"]})
        with self.assertRaises(ValidationError): save_client({"revision": 0, "name": "Новое"}, self.client1["id"])
        with self.assertRaises(ValidationError): save_rule({**self.rule, "amount": "0"}, self.rule["id"])

    def test_waiting_and_positive_intersection_same_employee_deterministic(self):
        booking = self.booking()
        self.assertEqual(booking["charge"]["state"], "waiting")
        boundary = self.work(start="10:00", end="12:00")
        self.assertEqual(ClientBonusCharge.objects.get().state, "waiting")
        row = self.work(start="12:59", end="14:00")
        row.refresh_from_db(); boundary.refresh_from_db()
        self.assertEqual((row.client_bonus, boundary.client_bonus), (200, 0))
        longer = self.work(start="12:00", end="13:00")
        row.refresh_from_db(); longer.refresh_from_db()
        self.assertEqual((row.client_bonus, longer.client_bonus), (0, 200))

    def test_choose_employee_and_stale_choice(self):
        row = self.work()
        other = Employee.objects.create(full_name="Other employee", short_name="Other employee")
        peer = self.work(employee=other)
        booking = self.booking()
        self.assertEqual(booking["charge"]["state"], "needs_choice")
        row.refresh_from_db(); self.assertTrue(row.client_bonus_review)
        query=f'?month=2026-04&organization={self.org.pk}'
        reviewed=self.client.get('/api/shifts/ledger/records/'+query+'&status=review').json()
        self.assertEqual(reviewed['count'],2)
        self.assertEqual(self.client.get('/api/shifts/ledger/records/'+query+'&status=ready').json()['count'],0)
        self.assertEqual(self.client.get('/api/shifts/ledger/records/'+query).json()['review'],2)
        choice = {"revision": booking["revision"], "token": booking["choices"]["token"], "record": row.pk, "rule": self.rule["id"]}
        chosen = save_booking(choice, booking["id"], "choose")
        row.refresh_from_db(); self.assertEqual(row.client_bonus, 200); self.assertFalse(row.client_bonus_review)
        save_record({"end_time": "17:00"}, peer)
        row.refresh_from_db(); self.assertTrue(row.client_bonus_review); self.assertEqual(row.client_bonus, 0)
        with self.assertRaises(ValidationError): save_booking({**choice, "revision": chosen["revision"]}, booking["id"], "choose")

    def test_rule_edits_preview_only_real_difference_and_stale_preview(self):
        row = self.work(); self.booking()
        changed = save_rule({**self.rule, "amount": "300"}, self.rule["id"])
        row.refresh_from_db(); self.assertEqual(row.client_bonus, 200)
        request = {"start": str(self.day), "end": str(self.day), "rule": changed["id"], "revision": changed["revision"]}
        preview = recalculate(request)
        self.assertEqual(len(preview["changes"]), 1)
        row.refresh_from_db(); self.assertEqual(row.client_bonus, 200)
        save_client({**self.client2, "name": "Изменено"}, self.client2["id"])
        with self.assertRaises(ValidationError): recalculate({**request, "fingerprint": preview["fingerprint"]})
        self.activate(changed)
        row.refresh_from_db(); self.assertEqual(row.client_bonus, 300)
        self.assertEqual(recalculate({"start": str(self.day), "end": str(self.day)})["changes"], [])

    def test_exclusion_restore_and_frozen_invoice_snapshot(self):
        row = self.work(); booking = self.booking()
        batch = prepare(self.day, self.day, [self.org.pk])
        inv = batch.invoices.get()
        from .models import OrganizationBilling, TelegramContact
        contact = TelegramContact.objects.create(user_id=900001, name="Test owner")
        OrganizationBilling.objects.get(organization=self.org).recipients.add(contact)
        from .billing import refresh
        refresh(batch); approve(batch.pk, batch.version)
        inv.refresh_from_db(); artifact = bytes(inv.artifact)
        with self.assertRaises(ValidationError): save_booking({"revision": booking["revision"], "reason": "Отмена"}, booking["id"], "exclude")
        excluded = save_booking({"revision": booking["revision"], "reason": "Отмена", "correction": True}, booking["id"], "exclude")
        row.refresh_from_db(); self.assertEqual(row.client_bonus, 0)
        inv.refresh_from_db(); self.assertEqual(bytes(inv.artifact), artifact)
        wb = load_workbook(BytesIO(artifact))
        text = str(list(wb['Основной'].values))
        self.assertIn("Школьные съёмки", text); self.assertNotIn("79191234567", text)
        restored = save_booking({"revision": excluded["revision"], "correction": True}, booking["id"], "restore")
        self.assertEqual(restored["state"], "draft")

    def test_staging_sanitization_deduplication_and_restart_lease(self):
        result = {"intervals": [{"room":"Первый", "start":"12:00", "end":"13:00",
                                "phone_candidates":[{"number":"+79191234567","score":.99}, {"number":"+79998887766","score":.99}]}]}
        sanitize_phones([result,result])
        self.assertEqual(result["intervals"][0]["clients"], [self.client1["id"]])
        self.assertNotIn("79191234567", json.dumps(result))
        imported = ClientBookingImport.objects.create(organization=self.org,date=self.day,images=[])
        job = RecognitionJob.objects.create(booking_import=imported,revision=1,available_at=timezone.now(),state="running",attempts=1)
        finish_import(job,[result]); imported.refresh_from_db()
        self.assertEqual(imported.state,"done")
        self.assertEqual(ClientBooking.objects.get().state,"draft")
        stage_bookings(result["intervals"],self.org.pk,self.day,{"type":"offer","id":99})
        self.assertEqual(ClientBooking.objects.count(),1)

    def test_bonus_does_not_share_limits_or_overlap_reduction(self):
        rate=Rate.objects.get(service=self.service,organization=self.org); rate.maximum=1000;rate.save()
        row=self.work(bonus="500");self.booking()
        row.refresh_from_db();self.assertEqual(row.amount,1700)
        save_record({"end_time":"17:00"},row)
        row.refresh_from_db();self.assertEqual((row.amount,row.client_bonus,row.bonus),(1700,200,500))

    def test_api_and_required_confirmation(self):
        row=self.work()
        payload={"client":self.client1["id"],"organization":self.org.pk,"date":str(self.day),"start_time":"12:00","end_time":"13:00"}
        response=self.client.post('/api/shifts/ledger/client-bookings/',json.dumps(payload),content_type='application/json')
        self.assertEqual(response.status_code,200)
        created=response.json();row.refresh_from_db();self.assertEqual(row.client_bonus,0)
        response=self.client.post(f'/api/shifts/ledger/client-bookings/{created["id"]}/confirm/',json.dumps({"revision":created["revision"]}),content_type='application/json')
        self.assertEqual(response.status_code,200)
        row.refresh_from_db();self.assertEqual(row.client_bonus,200)

    def test_multiple_rules_require_one_choice_and_exclusion_reverses_only_client_bonus(self):
        row=self.work(bonus="500")
        extra=save_rule({"name":"Другая доплата","organization":self.org.pk,"amount":"300","start":str(self.day),"clients":[self.client1["id"]]})
        self.activate(extra)
        booking=self.booking()
        row.refresh_from_db();self.assertTrue(row.client_bonus_review);self.assertEqual(row.client_bonus,0)
        selected=save_booking({"revision":booking["revision"],"token":booking["choices"]["token"],"rule":extra["id"],"record":row.pk},booking["id"],"choose")
        row.refresh_from_db();self.assertEqual((row.amount,row.client_bonus,row.bonus),(2400,300,500))
        save_booking({"revision":selected["revision"],"reason":"Съёмка отменена"},selected["id"],"exclude")
        row.refresh_from_db();self.assertEqual((row.amount,row.client_bonus,row.bonus),(2100,0,500))

    def test_ambiguous_card_selects_only_one_client(self):
        self.work()
        ids=stage_bookings([{"room":"Зал","start":"12:00","end":"13:00","clients":[self.client1["id"],self.client2["id"]],"phone_uncertain":True}],
                           self.org.pk,self.day,{"type":"import","id":99,"image":0})
        first,second=[ClientBooking.objects.get(pk=i) for i in ids]
        save_booking({"revision":first.revision},first.pk,"confirm")
        second.refresh_from_db();self.assertEqual(second.state,"excluded")
        with self.assertRaises(ValidationError):save_booking({"revision":second.revision},second.pk,"confirm")
        self.assertEqual(ClientBonusCharge.objects.filter(state="applied").count(),1)

    def test_partial_ocr_record_can_be_excluded_and_worker_stale_lease_cannot_commit(self):
        ids=stage_bookings([{"room":"Зал","start":"12:00","end":None,"clients":[self.client1["id"]]}],self.org.pk,self.day,{"type":"import","id":99})
        obj=ClientBooking.objects.get(pk=ids[0])
        excluded=save_booking({"revision":obj.revision,"reason":"Ошибочный снимок"},obj.pk,"exclude")
        self.assertEqual(excluded["state"],"excluded")
        imported=ClientBookingImport.objects.create(organization=self.org,date=self.day,images=[])
        job=RecognitionJob.objects.create(booking_import=imported,revision=1,available_at=timezone.now(),state="running",attempts=2)
        old=SimpleNamespace(pk=job.pk,booking_import_id=imported.pk,revision=1,attempts=1)
        finish_import(old,[]);imported.refresh_from_db();self.assertEqual(imported.state,"pending")

    def test_fixed_tariff_full_bonus_and_earnings_details_without_phones(self):
        rate=Rate.objects.get(service=self.service,organization=self.org);rate.calculation="fixed";rate.price=1200;rate.save()
        row=self.work(bonus="500");self.booking()
        row.refresh_from_db();self.assertEqual(row.amount,1900)
        from .earnings import report_data, workbook_bytes
        report=report_data("2026-04")
        self.assertEqual(report["client_bonuses"][0]["rule"],"Школьные съёмки")
        self.assertNotIn("79191234567",json.dumps(report,default=str))
        wb=load_workbook(BytesIO(workbook_bytes(report)))
        self.assertIn("Клиентские доплаты",wb.sheetnames)

    def test_rule_deactivation_is_applied_only_by_preview(self):
        row=self.work();self.booking()
        changed=save_rule({**self.rule,"active":False},self.rule["id"])
        self.assertFalse(changed["active"])
        row.refresh_from_db();self.assertEqual(row.client_bonus,200)
        self.activate(changed)
        row.refresh_from_db();self.assertEqual(row.client_bonus,0)

    def test_rule_with_disabled_client_can_be_switched_off(self):
        save_client({**self.client1,'active':False},self.client1['id'])
        changed=save_rule({**self.rule,'active':False},self.rule['id'])
        self.assertFalse(changed['active'])

    def test_telegram_edit_keeps_separate_client_and_manual_bonuses(self):
        from .ingest import ingest
        from .models import EmployeePreference
        self.employee.telegram_user_id=991;self.employee.save()
        EmployeePreference.objects.create(employee=self.employee,service=self.service)
        row=ingest('10:00-18:00 Bonus studio',chat_id=-100,message_id=991,user_id=991,today=self.day)[0]
        save_record({'bonus':'500','bonus_reason':'Дополнительная работа'},row)
        self.booking()
        edited=ingest('10:00-18:00 Bonus studio комментарий',chat_id=-100,message_id=991,user_id=991,today=self.day,edited=True)[0]
        self.assertEqual((edited.amount,edited.client_bonus,edited.bonus),(2300,200,500))

    def test_unchanged_rule_revision_is_not_a_financial_preview_change(self):
        row=self.work();self.booking()
        changed=save_rule(self.rule,self.rule['id'])
        request={'start':str(self.day),'end':str(self.day),'rule':changed['id'],'revision':changed['revision']}
        preview=recalculate(request);self.assertEqual(preview['changes'],[])
        recalculate({**request,'fingerprint':preview['fingerprint']})
        row.refresh_from_db();self.assertEqual(row.client_bonus,200)
        self.assertEqual(row.client_bonus_details[0]['rule_revision'],changed['revision'])

    def test_partial_studio_overlap_reduces_hours_but_keeps_both_bonuses_full(self):
        org=Organization.objects.create(name="Other overlapping studio")
        Rate.objects.create(service=self.service,organization=org,calculation="hourly",price=200,start=self.day)
        save_record({'kind':'service','service':self.service.pk,'organization':org.pk,'employee':self.employee.pk,
                     'date':str(self.day),'start_time':'10:00','end_time':'18:00'})
        rate=Rate.objects.get(organization=self.org,service=self.service);rate.minimum=1000;rate.save()
        row=self.work(start='16:00',end='20:00',bonus='500');self.booking(start='17:00',end='18:00')
        row.refresh_from_db()
        self.assertEqual((row.overlap_hours,row.payable_hours,row.amount,row.bonus,row.client_bonus),(2,3,1700,500,200))

    def test_moving_work_to_another_date_removes_client_charge(self):
        row=self.work();self.booking()
        save_record({'date':'2026-04-13'},row)
        row.refresh_from_db();self.assertEqual(row.client_bonus,0)
        self.assertEqual(ClientBonusCharge.objects.get().state,'waiting')
        save_record({'date':str(self.day)},row)
        row.refresh_from_db();self.assertEqual(row.client_bonus,200)

    def test_archived_employee_keeps_bonus_for_completed_work(self):
        row=self.work();self.booking()
        self.employee.is_active=False;self.employee.save()
        preview=recalculate({'start':str(self.day),'end':str(self.day)})
        self.assertEqual(preview['changes'],[])
        row.refresh_from_db();self.assertEqual(row.client_bonus,200)

    def test_phone_crop_does_not_read_the_next_booking(self):
        import numpy as np
        from apps.offers.recognizer import read_booking_phones
        intervals=[{"box":[0,.1,1,.9]},{"box":[0,.5,1,.9]}]
        class OCR:
            def read(self,crop):
                return [{"text":"8 (919) 123-45-67" if crop[0,0,0]<50 else "+79195555555","score":.99}]
        image=np.zeros((100,100,3),dtype=np.uint8);image[50:]=100
        read_booking_phones(intervals,image,OCR())
        self.assertEqual(intervals[0]["phone_candidates"][0]["number"],"+79191234567")
        self.assertEqual(intervals[1]["phone_candidates"][0]["number"],"+79195555555")

    def test_import_binary_upload_queue_confirmation_and_retry_revision(self):
        import tempfile
        from pathlib import Path
        from PIL import Image
        from django.core.files.uploadedfile import SimpleUploadedFile
        from django.test import override_settings
        from apps.offers.recognition_worker import run_once
        stream=BytesIO();Image.new("RGB",(200,300),"white").save(stream,format="PNG")
        root=Path(__file__).resolve().parents[3]/"tmp"
        with tempfile.TemporaryDirectory(dir=root) as directory, override_settings(OFFER_MEDIA_ROOT=Path(directory)):
            upload=SimpleUploadedFile("calendar.png",stream.getvalue(),content_type="image/png")
            response=self.client.post('/api/shifts/ledger/client-booking-imports/',{"organization":self.org.pk,"date":str(self.day),"images":[upload]})
            self.assertEqual(response.status_code,200)
            imported=response.json();url=f'/api/shifts/ledger/client-booking-imports/{imported["id"]}/'
            photo=self.client.get(url+'image/?index=0');self.assertEqual(photo.status_code,200)
            self.assertTrue(b''.join(photo.streaming_content).startswith(b'\x89PNG'))
            result={"date":str(self.day),"intervals":[{"room":"Зал","start":"12:00","end":"13:00","phone_candidates":[{"number":"+79191234567","score":.99}]}]}
            with patch('apps.offers.recognition_worker.analyze_image',return_value=result):
                self.assertTrue(run_once(object()))
            imported=self.client.get(url).json();self.assertEqual(imported["state"],"done")
            self.assertNotIn('79191234567',json.dumps(imported));self.assertEqual(ClientBooking.objects.get().state,'draft')
            obj=ClientBookingImport.objects.get(pk=imported["id"]);obj.state='failed';obj.save()
            response=self.client.post(url+'retry/',json.dumps({"revision":0}),content_type='application/json')
            self.assertEqual(response.status_code,400)
            response=self.client.post(url+'retry/',json.dumps({"revision":obj.revision}),content_type='application/json')
            self.assertEqual(response.status_code,200)
