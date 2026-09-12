from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
import re

from django.core.exceptions import ValidationError
from django.db import transaction
from django.forms.models import model_to_dict
from django.utils import timezone

from apps.shifts.audit import json_safe, log_change
from apps.shifts.models import AuditLog, Organization
from .engine import accrue, digest, lock, money, rate_for, record_dict
from .models import Accrual, Batch, Delivery, Invoice, OrganizationBilling, Rate, Record, Service


def safe_text(value):
    value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value or ""))[:32760]
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def workbook_bytes(invoice, batch):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    main = wb.active
    main.title = "Основной"
    main.append([safe_text(invoice.organization_name), f"{batch.start:%d.%m.%Y} — {batch.end:%d.%m.%Y}"])
    main.append(["Счёт", f"{invoice.pk} / версия {invoice.version}"])
    main.append(["Всего по счёту", float(invoice.total)])
    main.append([])
    columns = ["Дата", "Исполнитель", "Услуга / описание", "Время", "Часы / количество", "Сумма"]
    sheets = {"Основной": main}
    main.append(columns)
    totals = {}
    for line in invoice.lines:
        if line["id"] in invoice.excluded_ids:
            continue
        sheet = line["sheet"]
        if sheet not in sheets:
            sheets[sheet] = wb.create_sheet(sheet)
            sheets[sheet].append(columns)
        time_range = f'{line["start_time"][:5]}–{line["end_time"][:5]}' if line.get("start_time") and line.get("end_time") else ""
        title = line["description"] if line["kind"] in ("oneoff", "expense") else line["service_name"] + (f' — {line["description"]}' if line["description"] else "")
        sheets[sheet].append([date.fromisoformat(line["date"]), safe_text(line["employee_name"]), safe_text(title), time_range,
                              float(line["units"]) if line.get("input_type") in ("time", "quantity") else None, float(line["amount"])])
        totals[sheet] = totals.get(sheet, Decimal(0)) + Decimal(line["amount"])
    for name, ws in sheets.items():
        ws.append(["Итого по листу", None, None, None, None, float(totals.get(name, 0))])
    main.append([])
    main.append(["Доплаты и скидки"])
    for item in invoice.adjustments:
        main.append([safe_text(item["description"]), None, None, None, None, float(item["amount"])])
    main.append([])
    main.append(["Сводка по листам"])
    for name, total in totals.items():
        main.append([name, float(total)])
    main.append(["Корректировки", float(sum(Decimal(a["amount"]) for a in invoice.adjustments))])
    main.append(["Общий итог", float(invoice.total)])
    for ws in wb.worksheets:
        ws.freeze_panes = "A6" if ws == main else "A2"
        ws.sheet_view.showGridLines = False
        header_row = 5 if ws == main else 1
        for cell in ws[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="28756A")
        for index, width in enumerate([15, 25, 54, 20, 22, 20], 1):
            ws.column_dimensions[get_column_letter(index)].width = width
        for row in ws:
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, date):
                    cell.number_format = "DD.MM.YYYY"
                elif isinstance(cell.value, (float, int)):
                    cell.number_format = '#,##0.00'
        ws.print_options.horizontalCentered = True
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
    result = BytesIO()
    wb.save(result)
    return result.getvalue()


def current_source(batch):
    profiles = []
    for oid in batch.organization_ids:
        profile, _ = OrganizationBilling.objects.get_or_create(organization_id=oid)
        org = Organization.objects.get(pk=oid)
        profiles.append({"organization": oid, "name": org.name, "active": org.is_active,
                         "recipients": list(profile.recipients.order_by("user_id").values_list("user_id", flat=True))})
    records = [record_dict(r) for r in Record.objects.filter(organization_id__in=batch.organization_ids, date__range=(batch.start, batch.end), deleted_at=None).select_related("organization", "service").order_by("id")]
    return {"profiles": profiles, "records": records,
            "rates": list(Rate.objects.filter(organization_id__in=batch.organization_ids).order_by("id").values()),
            "services": list(Service.objects.order_by("id").values()),
            "accruals": list(Accrual.objects.filter(organization_id__in=batch.organization_ids).order_by("id").values())}


@transaction.atomic
def prepare(start, end, organizations, scheduled_key=None, replaces=None):
    config = lock()
    if start > end or end > timezone.localdate() or (end - start).days > 366:
        raise ValidationError("Выберите прошедший период не длиннее года.")
    organizations = sorted(set(int(x) for x in organizations))
    if not organizations or Organization.objects.filter(pk__in=organizations, is_active=True).count() != len(organizations):
        raise ValidationError("Выберите активные организации.")
    if scheduled_key:
        existing = Batch.objects.filter(scheduled_key=scheduled_key).first()
        if existing:
            return existing
    # Avoid accidentally sending an overlapping billing period twice.
    overlapping = Invoice.objects.filter(batch__state="approved", organization_id__in=organizations,
                                        batch__start__lte=end, batch__end__gte=start)
    if overlapping.exists() and not replaces:
        raise ValidationError("За пересекающийся период уже есть счёт. Используйте «Новая версия».")
    if replaces:
        previous = Batch.objects.get(pk=replaces)
        if previous.state != "approved" or previous.start != start or previous.end != end or previous.organization_ids != organizations:
            raise ValidationError("Новая версия должна заменять тот же период и организации.")
    accrue(min(end, timezone.localdate()))
    batch = Batch.objects.create(start=start, end=end, organization_ids=organizations, scheduled_key=scheduled_key, version=0,
                                 approver_id_snapshot=config.approver.user_id if config.approver_id else None)
    for oid in organizations:
        prev = Invoice.objects.filter(batch_id=replaces, organization_id=oid).first() if replaces else None
        Invoice.objects.create(batch=batch, organization_id=oid, organization_name=Organization.objects.get(pk=oid).name,
                               replaces=prev, adjustments=prev.adjustments if prev else [], excluded_ids=[])
    refresh(batch)
    return batch


@transaction.atomic
def refresh(batch):
    config = lock()
    batch.refresh_from_db()
    if batch.state == "approved":
        raise ValidationError("Отправленный пакет неизменяем. Создайте новую версию.")
    accrue(min(batch.end, timezone.localdate()))
    data = current_source(batch)
    batch.version += 1
    batch.fingerprint = digest(data)
    batch.approver_id_snapshot = config.approver.user_id if config.approver_id else None
    batch.save()
    Delivery.objects.filter(batch=batch, state="pending").update(state="cancelled")
    for invoice in batch.invoices.all():
        profile = next(p for p in data["profiles"] if p["organization"] == invoice.organization_id)
        invoice.organization_name = profile["name"]
        invoice.recipients = profile["recipients"]
        invoice.version = (invoice.replaces.version + 1) if invoice.replaces_id else batch.version
        # Older flags remain in the schema for compatibility, but new invoices account for every live record.
        invoice.excluded_ids = []
        invoice.lines = [r for r in data["records"] if r["organization"] == invoice.organization_id]
        order = {s["id"]: s["sort_order"] for s in data["services"]}
        invoice.lines.sort(key=lambda r: (0 if r["sheet"] == "Основной" else 2 if r["sheet"] == "Расходы" else 1, order.get(r["service"], 100), r["date"], r["id"]))
        invoice.errors = [f'Запись #{r["id"]}: {r["error"] or "нужно проверить исполнителя"}' for r in invoice.lines
                          if r["id"] not in invoice.excluded_ids and (r["review"] or r["error"])]
        for line in invoice.lines:
            if line["kind"] == "service" and line["id"] not in invoice.excluded_ids:
                try:
                    rate_for(line["service"], line["organization"], date.fromisoformat(line["date"]))
                except ValidationError as error:
                    invoice.errors.append(f'Запись #{line["id"]}: {" ".join(error.messages)}')
        if not profile["active"]:
            invoice.errors.append("Организация отключена.")
        invoice.total = sum((Decimal(r["amount"]) for r in invoice.lines if r["id"] not in invoice.excluded_ids), Decimal(0)) + sum((Decimal(a["amount"]) for a in invoice.adjustments), Decimal(0))
        if invoice.total < 0:
            invoice.errors.append("Итог счёта не может быть отрицательным.")
        if invoice.total and not invoice.recipients:
            invoice.errors.append("Не назначен получатель Telegram.")
        invoice.artifact = workbook_bytes(invoice, batch)
        invoice.save()
    if batch.approver_id_snapshot:
        for invoice in batch.invoices.all():
            Delivery.objects.get_or_create(key=f"preview:{batch.pk}:{batch.version}:{invoice.pk}", defaults={
                "batch": batch, "invoice": invoice, "purpose": "preview", "recipient": batch.approver_id_snapshot, "version": batch.version})
        Delivery.objects.get_or_create(key=f"approval:{batch.pk}:{batch.version}", defaults={
            "batch": batch, "purpose": "approval", "recipient": batch.approver_id_snapshot, "version": batch.version})
    log_change("batch", batch.pk, "refresh", actor="system", after={"version": batch.version, "fingerprint": batch.fingerprint})
    return batch


@transaction.atomic
def approve(batch_id, version, actor="web", telegram_user=None):
    config = lock()
    batch = Batch.objects.select_for_update().get(pk=batch_id)
    if telegram_user is not None and (not config.approver_id or config.approver.user_id != telegram_user or batch.approver_id_snapshot != telegram_user):
        raise ValidationError("Подтверждение доступно только назначенному человеку.")
    if batch.state == "approved":
        return {"approved": True, "already_approved": True}
    if batch.version != int(version):
        raise ValidationError("Кнопка относится к старой версии. Откройте новый пакет.")
    accrue(min(batch.end, timezone.localdate()))
    if digest(current_source(batch)) != batch.fingerprint:
        refresh(batch)
        return {"approved": False, "refreshed": True, "message": "Данные изменились. Проверьте обновлённые файлы и подтвердите заново."}
    errors = [e for inv in batch.invoices.all() for e in inv.errors]
    if errors:
        raise ValidationError(errors)
    for invoice in batch.invoices.all():
        other = Invoice.objects.filter(organization=invoice.organization, batch__state="approved", batch__start__lte=batch.end, batch__end__gte=batch.start).exclude(pk=invoice.replaces_id)
        # Older versions in the same replacement chain are permitted.
        ancestor = invoice.replaces
        ancestor_ids = []
        while ancestor:
            ancestor_ids.append(ancestor.pk)
            ancestor = ancestor.replaces
        if other.exclude(pk__in=ancestor_ids).exists():
            raise ValidationError("За этот период уже подтвердили другой счёт. Создайте новую версию.")
    batch.state, batch.approved_by, batch.approved_at = "approved", actor, timezone.now()
    batch.save()
    for invoice in batch.invoices.all():
        if not invoice.total:
            continue
        for recipient in invoice.recipients:
            Delivery.objects.get_or_create(key=f"owner:{invoice.pk}:{recipient}", defaults={
                "batch": batch, "invoice": invoice, "purpose": "owner", "recipient": recipient, "version": batch.version})
    log_change("batch", batch.pk, "approve", actor=actor, after={"version": batch.version})
    return {"approved": True}


@transaction.atomic
def schedule(now=None):
    config = lock()
    now = timezone.localtime(now or timezone.now())
    if not config.enabled or now.day < config.day or (now.day == config.day and (now.hour, now.minute) < (config.hour, config.minute)):
        return None
    end = now.date().replace(day=1) - timedelta(days=1)
    key = end.strftime("%Y-%m")
    if AuditLog.objects.filter(entity_type="batch", action="delete", after__scheduled_key=key).exists():
        return None
    existing = Batch.objects.filter(scheduled_key=key).first()
    if existing:
        return existing
    ids = list(OrganizationBilling.objects.filter(monthly=True, organization__is_active=True).values_list("organization_id", flat=True))
    if not ids:
        return None
    # Manually approved organizations are not billed again by the monthly scheduler.
    done = Invoice.objects.filter(batch__state="approved", batch__start__lte=end, batch__end__gte=end.replace(day=1)).values_list("organization_id", flat=True)
    ids = [oid for oid in ids if oid not in done]
    return prepare(end.replace(day=1), end, ids, scheduled_key=key) if ids else None


@transaction.atomic
def delete_batch(batch_id):
    lock()
    batch = Batch.objects.get(pk=batch_id)
    if Delivery.objects.filter(batch=batch, state="sending").exists():
        raise ValidationError("Счёт сейчас отправляется. Дождитесь результата доставки и повторите удаление.")
    ids = list(batch.invoices.values_list("id", flat=True))
    # Detach replacement references before removing protected documents and their bytes.
    Invoice.objects.filter(replaces_id__in=ids).update(replaces=None)
    Delivery.objects.filter(batch=batch).delete()
    batch.invoices.all().delete()
    log_change("batch", batch.pk, "delete", actor="web", after={"scheduled_key": batch.scheduled_key, "invoice_ids": ids})
    batch.delete()


@transaction.atomic
def delete_invoice(invoice_id):
    lock()
    invoice = Invoice.objects.select_related("batch").get(pk=invoice_id)
    batch = invoice.batch
    if Delivery.objects.filter(batch=batch, state="sending").exists():
        raise ValidationError("Пакет сейчас отправляется. Дождитесь результата доставки и повторите удаление.")
    if batch.invoices.count() == 1:
        delete_batch(batch.pk)
        return
    oid = invoice.organization_id
    Invoice.objects.filter(replaces=invoice).update(replaces=None)
    Delivery.objects.filter(invoice=invoice).delete()
    log_change("invoice", invoice.pk, "delete", actor="web", after={"batch_id": batch.pk, "organization_id": oid})
    invoice.delete()
    batch.organization_ids = [value for value in batch.organization_ids if value != oid]
    batch.save(update_fields=("organization_ids",))
    if batch.state == "draft":
        refresh(batch)
