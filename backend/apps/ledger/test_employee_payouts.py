import json
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta, timezone as dtz
from decimal import Decimal
from io import BytesIO
from threading import Barrier
from unittest import skipUnless
from unittest.mock import AsyncMock
from types import SimpleNamespace

from asgiref.sync import async_to_sync
from django.core.exceptions import ValidationError
from django.db import connection, connections, close_old_connections
from django.db.models import Sum
from django.test import TestCase, TransactionTestCase, override_settings
from openpyxl import load_workbook

from apps.shifts.models import Employee, Organization, AuditLog
from .models import EmployeePayout, EmployeePayoutLine, EmployeePayoutMessage, OrganizationReceipt, Record, Settings, TelegramContact
from .payouts import preview, confirm, cancel, record_payment, details
from .receipt_bot import ingest_receipt
from .payout_bot import register_message, claim, finish, update_once, navigation, day_preview, callback, existing_message
from .earnings import report_data, workbook_bytes
from .engine import save_record, delete_record


class Fixture:
    def setUp(self):
        self.e = Employee.objects.create(short_name="Выплаты тест", telegram_user_id=98822)
        self.o = Organization.objects.create(name="Выплаты студия", aliases=["выпст"])
        self.admin = TelegramContact.objects.create(user_id=98821, name="Главный")
        Settings.objects.update_or_create(pk=1, defaults={"approver": self.admin})
        self.row = Record.objects.create(kind="oneoff", employee=self.e, employee_name=self.e.display_name,
                  organization=self.o, date=date(2026, 10, 5), amount=2000, description="Смена с доплатами")

    def post(self, body, path=""):
        return self.client.post("/api/shifts/ledger/employee-payouts/" + path, json.dumps(body), content_type="application/json")


@override_settings(SHIFT_SYNC_AFTER_WRITE=False)
class EmployeePayoutTests(Fixture, TestCase):
    def test_group_individual_idempotency_cancel_and_reports(self):
        extra = Record.objects.create(kind="oneoff", employee=self.e, organization=self.o, date=self.row.date, amount=400, description="Уборка")
        draft = self.post({"action":"preview", "record":extra.pk}).json()
        paid = self.post({"action":"confirm", "token":draft["token"]})
        self.assertEqual(paid.status_code, 200, paid.content)
        self.assertTrue(self.post({"action":"confirm", "token":draft["token"]}).json()["already_paid"])
        day = preview({"employee":self.e.pk, "organization":self.o.pk, "date":str(self.row.date)})
        self.assertEqual(Decimal(day["amount"]), 2000)
        payout = confirm({"token":day["token"]})
        self.assertEqual(EmployeePayout.objects.count(), 2)
        report = report_data("2026-10", datetime(2026,10,15,tzinfo=dtz.utc))
        self.assertEqual((report["total"],report["received_total"],report["remaining_total"]),("2400.00","2400.00","0.00"))
        wb = load_workbook(BytesIO(workbook_bytes(report)))
        self.assertEqual(wb["Заработок"]["C6"].value, 2400)
        self.assertEqual(wb["Заработок"]["D6"].value, 0)
        self.assertEqual(wb["Выплаты"]["F2"].value, 400)
        self.assertEqual(self.post({"revision":1}, f'{payout["id"]}/cancel/').status_code,400)
        cancel(payout["id"], {"revision":1,"reason":"Ошибочная отметка"})
        self.row.refresh_from_db()
        self.assertEqual(Decimal(record_payment(self.row)["payout_remaining"]),2000)
        self.assertEqual(AuditLog.objects.filter(entity_type="employee_payout", action="cancel").count(),1)

    def test_snapshot_stale_review_and_no_cross_org_or_employee(self):
        d = preview({"record":self.row.pk})
        Record.objects.filter(pk=self.row.pk).update(amount=2200)
        with self.assertRaises(ValidationError): confirm({"token":d["token"]})
        Record.objects.filter(pk=self.row.pk).update(review=True)
        self.assertEqual(self.post({"action":"preview","record":self.row.pk}).status_code,400)
        Record.objects.filter(pk=self.row.pk).update(review=False)
        other = Employee.objects.create(short_name="Другой выплаты")
        Record.objects.create(kind="oneoff", employee=other, organization=self.o, date=self.row.date, amount=900)
        day = preview({"employee":self.e.pk,"organization":self.o.pk,"date":str(self.row.date)})
        self.assertEqual(Decimal(day["amount"]),2200)
        Record.objects.create(kind="oneoff", employee=self.e, organization=self.o, date=self.row.date, amount=50)
        with self.assertRaises(ValidationError): confirm({"token":day["token"]})

    def test_paid_amount_survives_reprice_and_guards_identity_delete(self):
        d = preview({"record":self.row.pk});confirm({"token":d["token"]})
        changed = save_record({"amount":"2500"}, self.row)
        self.assertEqual(Decimal(record_payment(changed)["payout_remaining"]),500)
        d2 = preview({"record":self.row.pk});confirm({"token":d2["token"]})
        changed = save_record({"amount":"1900"}, self.row)
        self.assertEqual(record_payment(changed)["payout_state"], "overpaid")
        self.assertEqual(Decimal(record_payment(changed)["payout_remaining"]), -600)
        with self.assertRaises(ValidationError): save_record({"date":"2026-10-06"}, self.row)
        with self.assertRaises(ValidationError): delete_record(self.row)
        self.assertEqual(EmployeePayoutLine.objects.aggregate(total=Sum("amount"))["total"],2500)

    def test_tracking_cutoff_details_filters_authorization(self):
        Record.objects.filter(pk=self.row.pk).update(date=date(2026,9,30))
        self.assertEqual(self.post({"action":"preview","record":self.row.pk}).status_code,400)
        report = report_data("2026-09", datetime(2026,10,15,tzinfo=dtz.utc))
        self.assertIsNone(report["employees"][0]["received"])
        self.assertIsNone(report["remaining_total"])
        self.assertEqual(load_workbook(BytesIO(workbook_bytes(report)))["Заработок"]["C6"].value,"Нет данных")
        with self.assertRaises(ValidationError): preview({"record":self.row.pk}, actor="telegram:98822",user_id=98822)
        Record.objects.filter(pk=self.row.pk).update(date=date(2026,10,5))
        d = preview({"record":self.row.pk}, actor="telegram:98821",user_id=98821)
        with self.assertRaises(ValidationError): confirm({"token":d["token"]}, actor="web")
        with self.assertRaises(ValidationError): confirm({"token":d["token"]}, actor="telegram:98822",user_id=98822)
        confirm({"token":d["token"]}, actor="telegram:98821",user_id=98821)
        self.assertEqual(details({"month":"2026-10","employee":self.e.pk,"state":"paid"})["count"],1)
        self.assertEqual(details({"month":"2026-10","employee":self.e.pk,"state":"unpaid"})["count"],0)

    def test_message_lease_revision_restart_and_repeat_edit(self):
        item = register_message(-100, 22, [self.row.pk])
        job = claim()
        d = preview({"record":self.row.pk});confirm({"token":d["token"]})
        finish(job[0],job[1],job[2])
        item.refresh_from_db()
        self.assertEqual(item.state,"pending")
        self.assertEqual(item.revision,2)
        with self.assertRaises(ValidationError): day_preview(item.pk,1,self.row.pk,98821,-100,22)
        bot = AsyncMock()
        async_to_sync(update_once)(bot)
        self.assertIn("выплачено 2000",bot.edit_message_text.call_args.kwargs["text"])
        self.assertIsNone(bot.edit_message_text.call_args.kwargs["reply_markup"])
        self.assertFalse(async_to_sync(update_once)(bot))
        text,markup = navigation(98821)
        self.assertTrue(any(b.callback_data.startswith("payout:months:") for row in markup.inline_keyboard for b in row))

    def test_new_service_restores_button_and_preserves_other_payment(self):
        item=register_message(-100,22,[self.row.pk])
        draft=preview({"record":self.row.pk});result=confirm({"token":draft["token"]})
        bot=AsyncMock();async_to_sync(update_once)(bot)
        save_record({"kind":"oneoff","employee":self.e.pk,"organization":self.o.pk,"date":str(self.row.date),"amount":"100","description":"Доп услуга"})
        item.refresh_from_db();self.assertEqual(item.state,"pending")
        async_to_sync(update_once)(bot)
        self.assertIsNotNone(bot.edit_message_text.call_args.kwargs["reply_markup"])
        self.assertIn("осталось 100",bot.edit_message_text.call_args.kwargs["text"])

    def test_worker_restart_reclaims_edits_and_old_lease_cannot_finish(self):
        from django.utils import timezone
        item=register_message(-100,22,[self.row.pk],11)
        self.assertTrue(existing_message(-100,11))
        old=claim()
        EmployeePayoutMessage.objects.filter(pk=item.pk).update(started_at=timezone.now()-timedelta(minutes=6))
        new=claim();self.assertNotEqual(old[1],new[1])
        finish(old[0],old[1],old[2]);item.refresh_from_db();self.assertEqual(item.state,"sending")
        finish(new[0],new[1],new[2]);item.refresh_from_db();self.assertEqual(item.state,"sent")

    def test_callback_requires_admin_before_preview_and_confirm(self):
        item=register_message(-100,22,[self.row.pk])
        query=SimpleNamespace(data=f'payout:day:{item.pk}:1:{self.row.pk}',from_user=SimpleNamespace(id=98822),
            message=SimpleNamespace(chat=SimpleNamespace(id=-100,type="supergroup"),message_id=22),answer=AsyncMock())
        bot=AsyncMock();async_to_sync(callback)(query,bot)
        self.assertTrue(query.answer.call_args.kwargs["show_alert"]);self.assertEqual(EmployeePayout.objects.count(),0)
        query.from_user.id=98821;query.answer.reset_mock();async_to_sync(callback)(query,bot)
        markup=bot.send_message.call_args.kwargs["reply_markup"]
        query.data=markup.inline_keyboard[0][0].callback_data
        async_to_sync(callback)(query,bot);async_to_sync(callback)(query,bot)
        self.assertEqual(EmployeePayout.objects.count(),1)

    def test_late_old_telegram_edit_requeues_current_content(self):
        from django.utils import timezone
        item=register_message(-100,22,[self.row.pk])
        old=claim()
        draft=preview({"record":self.row.pk});confirm({"token":draft["token"]})
        EmployeePayoutMessage.objects.filter(pk=item.pk).update(started_at=timezone.now()-timedelta(minutes=6))
        new=claim();finish(new[0],new[1],new[2])
        finish(old[0],old[1],old[2])
        item.refresh_from_db();self.assertEqual(item.state,"pending")
        final=claim();self.assertIsNone(final[-1]);self.assertIn("выплачено 2000",final[-2])

    def test_owner_invoice_unchanged_by_payouts(self):
        from .billing import prepare,approve
        from .models import OrganizationBilling
        profile,_=OrganizationBilling.objects.get_or_create(organization=self.o)
        profile.recipients.add(self.admin)
        batch=prepare(self.row.date,self.row.date,[self.o.pk]);invoice=batch.invoices.first()
        approve(batch.pk,batch.version)
        artifact=bytes(invoice.artifact);fingerprint=batch.fingerprint
        d=preview({"record":self.row.pk});payment=confirm({"token":d["token"]})
        invoice.refresh_from_db();batch.refresh_from_db()
        self.assertEqual(bytes(invoice.artifact),artifact);self.assertEqual(batch.fingerprint,fingerprint)
        cancel(payment["id"],{"revision":1,"reason":"Тест"})
        invoice.refresh_from_db();self.assertEqual(invoice.total,2000)

    def test_paid_work_needing_review_is_not_reported_as_false_overpayment(self):
        draft=preview({"record":self.row.pk});confirm({"token":draft["token"]})
        Record.objects.filter(pk=self.row.pk).update(calculation_review=True)
        report=report_data("2026-10",datetime(2026,10,15,tzinfo=dtz.utc))
        self.assertTrue(report["payout_review"])
        self.assertEqual(report["received_total"],"2000.00")
        self.assertIsNone(report["remaining_total"])
        self.assertEqual(load_workbook(BytesIO(workbook_bytes(report)))["Заработок"]["D6"].value,"На проверке")


class ReceiptTelegramTests(Fixture, TestCase):
    def receipt(self,text,**kwargs):
        return ingest_receipt(text,user_id=kwargs.pop("user_id",98822),chat_id=12,message_id=kwargs.pop("message_id",1),today=date(2026,10,5),**kwargs)

    def test_receipt_deduction_aliases_dates_duplicates_and_author(self):
        first=self.receipt("получено 12.10 Анна 5 000,50 выпст")
        self.assertEqual(first["amount"],"5000.50")
        self.assertEqual(self.receipt("получено 12.10 Анна 5 000,50 выпст")["id"],first["id"])
        self.receipt("вычет 13.10.2026 Возврат клиенту 1000 выпст",message_id=2,user_id=98821)
        result=self.client.get(f'/api/shifts/ledger/receipts/?organization={self.o.pk}').json()
        self.assertEqual(Decimal(result["total"]),Decimal("4000.50"))
        self.assertEqual(OrganizationReceipt.objects.count(),2)
        self.assertEqual(Record.objects.get().amount,2000)
        self.assertEqual(AuditLog.objects.filter(entity_type="receipt",actor="telegram:98822").count(),1)

    def test_edit_validation_site_conflict_delete_and_rights(self):
        first=self.receipt("получено Анна 5000 выпст")
        changed=self.receipt("получено Анна 6000 выпст",edited=True)
        self.assertEqual(Decimal(changed["amount"]),6000)
        with self.assertRaises(ValidationError): self.receipt("вычет 0 выпст",edited=True)
        with self.assertRaises(ValidationError): self.receipt("получено Анна 6000 выпст",user_id=12345,message_id=3)
        url=f'/api/shifts/ledger/receipts/{first["id"]}/'
        response=self.client.put(url,json.dumps({**changed,"amount":"7000"}),content_type="application/json")
        self.assertEqual(response.status_code,200,response.content)
        with self.assertRaises(ValidationError): self.receipt("получено Анна 8000 выпст",edited=True)
        self.assertEqual(OrganizationReceipt.objects.get().amount,7000)
        removed=self.client.delete(url,json.dumps(response.json()),content_type="application/json")
        self.assertEqual(removed.status_code,200)
        with self.assertRaises(ValidationError): self.receipt("получено Анна 5000 выпст")

    def test_old_edit_delivery_is_ignored(self):
        self.receipt("получено Анна 5000 выпст",event_time=datetime(2026,10,5,1,tzinfo=dtz.utc))
        self.receipt("получено Анна 7000 выпст",edited=True,event_time=datetime(2026,10,5,3,tzinfo=dtz.utc))
        self.receipt("получено Анна 6000 выпст",edited=True,event_time=datetime(2026,10,5,2,tzinfo=dtz.utc))
        self.assertEqual(OrganizationReceipt.objects.get().amount,7000)

    def test_private_and_service_topic_receipt_messages(self):
        from apps.shifts.management.commands.run_shift_bot import process_receipt
        Settings.objects.filter(pk=1).update(service_chat=-100,service_thread=10)
        message=SimpleNamespace(text="получено Анна 5000 выпст",from_user=SimpleNamespace(id=98822),
            chat=SimpleNamespace(id=12,type="private"),message_id=1,date=datetime(2026,10,5,20,tzinfo=dtz.utc),edit_date=None,reply=AsyncMock())
        self.assertTrue(async_to_sync(process_receipt)(message))
        self.assertIn("5000",message.reply.call_args.args[0]);self.assertEqual(OrganizationReceipt.objects.get().date,date(2026,10,6))
        message.chat=SimpleNamespace(id=-100,type="supergroup");message.message_thread_id=10;message.message_id=2
        self.assertTrue(async_to_sync(process_receipt)(message))
        message.message_thread_id=20;message.message_id=3
        self.assertFalse(async_to_sync(process_receipt)(message));self.assertEqual(OrganizationReceipt.objects.count(),2)


@skipUnless(connection.vendor=="postgresql","PostgreSQL row locking")
class PayoutConcurrencyTests(Fixture, TransactionTestCase):
    def test_web_and_telegram_only_one_payment(self):
        web=preview({"record":self.row.pk})
        tg=preview({"record":self.row.pk},actor="telegram:98821",user_id=98821)
        barrier=Barrier(2)
        def run(args):
            close_old_connections();barrier.wait()
            try:
                confirm({"token":args[0]},**args[1]);return "paid"
            except ValidationError:return "stale"
            finally:connections["default"].close()
        with ThreadPoolExecutor(2) as pool:
            results=list(pool.map(run,[(web["token"],{}),(tg["token"],{"actor":"telegram:98821","user_id":98821})]))
        self.assertCountEqual(results,["paid","stale"])
        self.assertEqual(EmployeePayout.objects.count(),1)
