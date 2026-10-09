from datetime import date
from decimal import Decimal
from io import BytesIO
from django.core.exceptions import ValidationError
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from .test_overlap_receipts import OverlapAndReceiptTests
from .models import Rate, Record
from .engine import save_record, delete_record
from .overlaps import review_day, decide
from .billing import prepare


@override_settings(SHARED_SHIFT_ALLOCATION_ENABLED=False, SHIFT_SYNC_AFTER_WRITE=False, TELEGRAM_BOT_TOKEN="")
class OverlappingHoursTests(TestCase):
    setUp = OverlapAndReceiptTests.setUp
    shift = OverlapAndReceiptTests.shift

    def test_partial_union_and_bonus(self):
        Rate.objects.all().update(minimum=None, maximum=None)
        self.shift(0, "10:00", "18:00")
        b = self.shift(1, "16:00", "20:00", bonus="500", bonus_reason="Выход")
        self.assertEqual((b.actual_hours, b.overlap_hours, b.payable_hours, b.amount), (4, 2, 3, 1100))
        self.shift(0, "17:00", "18:00")
        b.refresh_from_db()
        self.assertEqual((b.overlap_hours, b.payable_hours, b.amount), (2, 3, 1100))

    def test_minimum_maximum_after_weighted_hours_and_fixed(self):
        Rate.objects.all().update(minimum=1000, maximum=None)
        self.shift(0, "10:00", "18:00")
        b = self.shift(1, "16:00", "20:00")
        self.assertEqual((b.payable_hours, b.amount), (3, 1000))
        Rate.objects.filter(organization=self.orgs[1]).update(minimum=None, maximum=500)
        b=save_record({"end_time":"20:01"}, b)
        self.assertEqual(b.amount, 500)
        Rate.objects.filter(organization=self.orgs[1]).update(calculation="fixed", price=1250, maximum=None)
        b=save_record({"end_time":"20:02"}, b)
        self.assertEqual((b.amount,b.overlap_factor,b.payable_hours), (1250,1,b.actual_hours))

    def test_minute_overlap_rounding_and_touching(self):
        Rate.objects.all().update(minimum=None, maximum=None)
        self.shift(0, "10:00", "18:00")
        b=self.shift(1, "17:59", "20:00")
        self.assertEqual(b.amount, Decimal("401.67"))
        self.assertEqual(b.overlap_hours, Decimal("0.016667"))
        b=save_record({"start_time":"18:00"}, b)
        self.assertEqual((b.amount,b.overlap_hours), (400,0))

    def test_manual_one_both_neither_and_stale_decision(self):
        Rate.objects.all().update(minimum=None, maximum=None)
        a=self.shift(0,"10:00","18:00");b=self.shift(1,"16:00","20:00")
        for selected, amounts in [([self.orgs[0].pk],(1400,800)),([o.pk for o in self.orgs[:2]],(1400,600)),([],(1600,800))]:
            component=review_day(self.employee.pk,self.day)[0]
            decide({**component,"half_organizations":selected},actor="Тест")
            a.refresh_from_db();b.refresh_from_db();self.assertEqual((a.amount,b.amount),amounts)
        stale={**component,"half_organizations":[]}
        with self.assertRaises(ValidationError): decide(stale)
        Rate.objects.filter(organization=self.orgs[1]).update(maximum=3000)
        # Even a tariff change not changing the sum invalidates the decision.
        current=review_day(self.employee.pk,self.day)[0]
        self.assertTrue(current["requires_review"])
        invoice=prepare(self.day,self.day,[self.orgs[1].pk]).invoices.get()
        self.assertTrue(any("ручной" in e for e in invoice.errors))
        decide({**current,"half_organizations":[]})
        a=save_record({"end_time":"17:00"},a)
        self.assertTrue(a.calculation_review)

    def test_three_studios_union_does_not_double_discount(self):
        Rate.objects.all().update(minimum=None, maximum=None)
        rows=[self.shift(0,"10:00","18:00"),self.shift(1,"12:00","16:00"),self.shift(2,"13:00","17:00")]
        component=review_day(self.employee.pk,self.day)[0]
        decide({**component,"half_organizations":[self.orgs[0].pk]})
        rows[0].refresh_from_db()
        self.assertEqual((rows[0].overlap_hours,rows[0].payable_hours,rows[0].amount),(5,Decimal("5.5"),1100))

    def test_excel_contains_weighted_hours_and_full_bonus(self):
        Rate.objects.all().update(minimum=None, maximum=None)
        self.shift(0,"10:00","18:00")
        b=self.shift(1,"16:00","20:00",bonus="500",bonus_reason="Выход")
        invoice=prepare(self.day,self.day,[self.orgs[1].pk]).invoices.get()
        ws=load_workbook(BytesIO(bytes(invoice.artifact)))["Основной"]
        self.assertEqual([ws.cell(6,c).value for c in (5,6,7,8,9,10)],[4,1100,800,2,3,500])
